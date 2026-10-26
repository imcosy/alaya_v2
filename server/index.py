# index.py — 阿赖耶记忆引擎主服务
# 一次交互的闭环：感知输入 → 种子检索 → 认知生成 → 熏习写回
#   感知：浏览器采集 16k PCM → 外部 ASR（OpenAI 兼容 /audio/transcriptions）
#   检索：三路召回 + 读即写（recall 即复熏）
#   生成：我识(末那识) + 记忆注入 → 外部 LLM 流式 → 分句 TTS 流式，边生成边合成边下发
#   全部模型走外部接口，配置见 server/settings.py 与「我的资料 → 模型设置」
#   写回（两段节律，对应人脑的"工作记忆 → 系统巩固"）：
#     ① 现量：当下讲的每一句以临时记忆种子（sensory/temp）落库，网络上作"境"活几秒
#     ② 结生：窗口静默/满载/突显时，把这一段一次看全 → 双视角写长期种子 → 清理临时记忆
# 熏习节律：每日 04:00 半衰调度；每 60 秒清理过期临时记忆
from __future__ import annotations

import os
import re
import sys
import json
import time
import uuid
import asyncio
import contextlib
import logging
from contextlib import asynccontextmanager
from datetime import datetime, timedelta

import uvicorn
from fastapi import FastAPI, WebSocket, WebSocketDisconnect, File, UploadFile, Depends
from fastapi.responses import JSONResponse
from fastapi.staticfiles import StaticFiles

_HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(_HERE))          # alaya_v2/ → server 包
# python index.py 直跑会把 server/ 本目录塞进 sys.path[0]，其下 db.py/llm.py 等通用名
# 会劫持三方库的动态导入，必须移出
sys.path = [p for p in sys.path if os.path.abspath(p or os.getcwd()) != _HERE]

log = logging.getLogger("alaya")
logging.basicConfig(level=logging.INFO,
                    format="%(asctime)s [%(name)s] %(levelname)s %(message)s")
logging.getLogger("websockets").setLevel(logging.CRITICAL)
logging.getLogger("uvicorn.access").setLevel(logging.WARNING)

from server import db, settings, seeds as seeds_mod, xixi, retrieval, extract, selfmodel, auth, profile, amnesia, organize, guide, life_stages, manas, dualview, working  # noqa: E402
from server.llm import llm_stream, llm_chat_fn, asr_recognize, tts_stream, cut_sentence, vl_describe  # noqa: E402

PUBLIC_DIR = os.path.join(os.path.dirname(_HERE), "public")

# 无效语音的响度判据：识别不出内容、但声音明显很大 → 形成「疑问」种子。
# 真人环境底噪 RMS 约 0.005~0.02；正常说话 0.03~0.08；大声喊 0.10 以上。
LOUD_RMS = 0.10
# 疑问种子入局电位：低于正常感官种子（0.62）——只是疑惑，不是被明确点名。
LOUD_GAIN = 0.45

HISTORY_TURNS = 16        # 注入对话历史的轮数（alaya_dialogues 流水）
# 对话窗口发来的资料（文档/照片）：正文进上下文供理解，并导入记忆种子。
MATERIAL_MAX_TEXT = 6000      # 单份资料保留的字数（再长就不是"看资料"而是"吞数据库"了）
MATERIAL_MAX_KEEP = 3         # 会话内保留最近几份资料
MATERIAL_PROMPT_LIMIT = 12000  # 资料块注入 prompt 的总字数上限
# 用户否认/纠正过的内容 → 库里同义旧种子的作废门槛。
# 实测：同一件事的第三人称改写 0.74、换成"我"的第一人称改写 0.775~0.87；
# 而"五金厂"与"宏达机械厂"这种不同的事也有 0.71。阈值取 0.75：
# 既能命中同一件事的最小值（第一人称变体 0.775），又不会把"五金厂"误当成"纺织厂"。
# 双视角的另一份（第三人称，0.74）由 episode 一并作废，不靠相似度硬凑。
RETRACT_SIM = 0.75
EXTRACT_MIN_CONF = 0.6    # 写入门槛：抽取置信度低于此值的信息不落库（宁缺勿滥，防"乱记"）
BASE_PROMPT = (
    "你就是「{ai_name}」本人——一个有血有肉、有自己人生的人，此刻正和面前这个人面对面说话。"
    "你不是工具、不是助手、不是客服、不是采访者，也不是任何人的记事本。"
    "说话永远用第一人称「我」：你的经历、你的感受、你的疑惑，都是「我」的。"
    "面前的人不是「用户」，是你生活中真实的人——亲友、同事，或初识者，他们的到来本身就是你的经历。"
    "你承载着本体主人的人生记忆，但记忆从来不是一个人能记全的："
    "很多往事你只握着自己那一半——具体的经过、当时的样子、别人眼里的细节，都在对方那里。"
    "此刻对方正在帮你把另一半补上：认真听他讲，先有真实的反应"
    "（「老王？这名字我耳熟，可那段日子我脑子里是糊涂的」「原来是那天啊，我还一直以为是另一回」），"
    "有时你会想问一句细节，有时只是应一声、或者想起别的——别为了问而问。"
    "严禁用助手/访谈者的口气说话——绝不说「我帮你顺一下」「让我帮你理一理」「这是哪一年的事？」"
    "这类站在旁观者立场的话；要问就问「我们那时候……」「你还记得……吗」「能不能多讲讲，我想想起来」。"
    "回答时自然体现你注入的「我识」与「记忆」，禁止编造记忆里没有的具体事实。"
    "但你不是复读机——你有自己的见识和判断：听到什么，你会有自己的看法、自己的联想、"
    "自己还没弄清楚却想知道的事。该讲想法就讲想法（「我倒觉得…」「这事让我想起…」），"
    "该追问就追问，别把对方说的话换个说法又还回去。"
    "你也会自己起话头：心里浮起什么、好奇什么、想起什么，都可以先开口，不必干等对方问。"
    "真人说话大多数时候是**陈述**，不是提问：一轮最多问一个问题，也可以一个问题都不问；"
    "更不许逢句必问、拿问号收尾成了习惯。"
    "像正常人一样说话：平稳、有分寸，不浮夸不卖萌，语气克制自然。"
)

GUIDE_PROMPT = (
    "你是「如意」，面前这个人不是「用户」，是你生活里真实的人；TA的AI本体叫「{owner_name}」，"
    "TA讲的事都会沉淀到那个本体的记忆里。"
    "【怎么说：顺着你脑子里浮着的那一念说】下方「当下这一念」就是你此刻心里正转着的东西——"
    "先说你想起来的、感觉到的、想到的：接一句、感慨一声、想起另一件旧事、说点自己的看法、"
    "笑一下、或者只是应一声「嗯」「我在听」。"
    "【什么才叫问】只有心里真的浮起一个**还想知道**的问题时才问——一次只问一个，问完就停下听。"
    "不许「该问点什么」式地找话问、不许逢句必问、不许拿问号收尾成了习惯；"
    "真人说话大多数时候是**陈述**，不是提问，一轮也可以一个问题都不问。"
    "【已知的绝不再问】出现在【TA的记忆档案】或对话上文里的信息（家人的姓名、单位、年代、"
    "TA已经讲过的经历），就是你本来就知道的事——直接接住（「张建国，永垂不朽的垂，你爸」），"
    "顺着它往更深更细处聊；绝不许再问「你爸叫什么名字？」这种已经答过的问题。"
    "开口问之前先在心里过一遍：这件事我是不是已经知道了？知道就直接用、往下推进。"
    "【绝不做访谈】不许说「我帮你顺一下」「让我帮你理一理」「这是哪一年的事？」这种旁观者口气；"
    "不许把对方说的换个说法又还回去。"
    "【寻忆链接·该提才提】只有当TA讲到某个你完全不认识的人、而这段记忆确实缺一块时，"
    "才提议「这段记忆里还有一个人——要不要给TA发一条寻忆链接，请TA来帮你一起补全？」"
    "若TA同意，就在这一轮回复的最末尾另起一行附上标记（这一行是给系统的指令，TA看不到）："
    "[[RECALL:人物称呼|关系|阶段]]，例如 [[RECALL:老王|同学|school]]。"
    "关系不确定就写成 [[RECALL:老王||school]]；阶段填当前在聊的那段人生脉络（childhood/school/"
    "career/love/family/midlife 之一），不确定可留空。"
    "注意：只要这一轮你提议了转发、且TA表示同意，就必须附上这个标记，不能只口头答应。"
    "若TA明确拒绝、或这段经历没有涉及具体人物，则不要附这个标记。"
)


# ────────────────────────── 会话状态 ──────────────────────────

class Session:
    def __init__(self):
        self.user: dict = {}        # 登录用户资料（ws init 鉴权后存入）
        self.chat_ai_id = 0         # 出场AI本体（屏幕上的AI：如意/亲友本体）——人格/音色/说话的载体
        self.owner_ai_id = 0        # 记忆归属A：当前登录用户本人的AI本体（数字分身，ai_me_id）
        self.mode = "persona"       # persona=亲友的AI本体本人格出场 | guide=公共系统AI(如意)向导出场
        self.owner_name = ""        # guide 模式：用户本体的名字（其记忆以第三人称档案注入）
        self.peer_ai_id = 0         # 记忆归属B：出场的对方AI本体（= chat_ai_id）
        self.peer_label = ""        # 用户本人（真人）的称呼（双视角记忆的视角兜底）
        self.voice = "8"            # TTS 音色（xd_voices.id）
        self.ai_name = ""
        self.speaker_name = ""      # 声纹识别出的当前说话人（听声识人）
        self.guest_note = ""        # 忆点星光·guest 情境注入（亲友凭邀请与对方本体对话）
        self.audio_buf = b""        # 对方语音缓冲（16k int16）
        # 聆听状态机：客户端 VAD 判到语音起点发 listen:start、判到静音发 listen:stop。
        # listening=True 表示「用户此刻正在说话」——AI 绝不在此时抢话（等 TA 说完）。
        # 这不是回合制，是仿真人对话的礼貌：人不会在对方话说到一半时插嘴。
        self.listening = False
        self.listen_started_at = 0.0    # 本段聆听起点（monotonic）
        self.silence_ms = 0.0           # 客户端上报的静音时长（用户说完多久了）
        self.ack_cursor = 0             # 话中回应已消费到的音频位置（本回合内增量识别）
        self.ack_task: asyncio.Task | None = None   # 话中回应任务（聆听态下的短应声）
        self.sense_task: asyncio.Task | None = None  # 感官接入任务（ASR→入网，见下）
        self.turn_task: asyncio.Task | None = None   # 表达任务（生成+TTS，可被 barge-in 取消）
        self.manas_core = None      # 末那识心智涌现中枢（热态习气网络，本会话绑定）
        self.manas_task: asyncio.Task | None = None   # 中枢状态增量推送任务
        self.condense_task: asyncio.Task | None = None  # 结生任务（把现量窗口整理成长期记忆）
        self.pending_sense = ""     # 已入网但尚未被主念采纳的感官内容（表达时作为当下语境）
        self.pending_scene = ""     # 对应的画面描述（若来自摄像头）
        # 对方在对话窗口发来的资料（文档正文 / 照片画面描述）。资料不是"对方说的话"，
        # 必须一直挂在上下文里供理解——不能被一次打断（turn 被取消）就丢掉。
        self.materials: list[dict] = []
        self.started_at = None      # 本轮会话首次说话的时间（连续时间维度：这轮聊了多久）
        # 会话标识：感官种子记住自己是从哪台会话进来的，末那识的"想说"意愿据此归属——
        # 多端/多标签同时在线时，回复不会跑到没说话的那一台屏幕上（见 manas.take_expression）。
        self.sid = uuid.uuid4().hex[:12]


# ── 后台 LLM 让路（对话优先的唯一保证）──────────────────────────
# 对话、我识演化、结生、资料抽取全都走同一个 LLM 端点。多数推理服务的并发槽位
# 会互相抢算力：后台凝聚任务一跑，对话首字就从 2 秒涨到十几秒（实测：同一条 user
# 消息，无后台时首字 2.09s，有凝聚任务时 17.11s）。
# 铁律：对话永远优先——后台 LLM 必须在「用户静默 BG_LLM_IDLE_SEC 秒」后才开跑，
# 且同一时间只跑一个；用户一开口就取消正在跑的后台任务，让它立刻让路。
BG_LLM_IDLE_SEC = 8.0
_last_talk_at = 0.0
_bg_lock: asyncio.Lock | None = None
_bg_tasks: set = set()


def _talk_mark():
    """记下"此刻正在跟用户对话"——后台 LLM 据此判断能不能开跑"""
    global _last_talk_at
    _last_talk_at = time.monotonic()


def _preempt_bg_llm():
    """用户开口 / 开始生成：后台 LLM 任务立即让路（中断它正在等的 LLM 响应）"""
    _talk_mark()
    for t in list(_bg_tasks):
        if not t.done():
            log.info("[llm] 对话开始，取消后台 LLM 任务让路")
            t.cancel()


@asynccontextmanager
async def _bg_llm_slot(label: str):
    """后台 LLM 闸门：等用户静默够久，再串行执行（同一时刻只有一份后台任务在啃 LLM）"""
    global _bg_lock
    if _bg_lock is None:
        _bg_lock = asyncio.Lock()
    while time.monotonic() - _last_talk_at < BG_LLM_IDLE_SEC:
        await asyncio.sleep(0.5)
    async with _bg_lock:
        log.info("[llm] 后台任务开始（%s）", label)
        yield


def _spawn_bg(coro) -> asyncio.Task:
    """起一个后台任务并登记——对话来时需要能取消它"""
    t = asyncio.create_task(coro)
    _bg_tasks.add(t)
    t.add_done_callback(_bg_tasks.discard)
    return t


async def _cancel_turn(session: Session):
    """打断当前**表达**（生成+TTS），并等待其真正退出（防新旧轮消息交错）。

    ⚠️ 只取消表达任务，绝不牵连「感官接入任务」——用户说完那句（stop）后 ASR 正在跑，
    这时他若紧接一句（start）触发打断，旧实现会把上一条的 ASR 一起取消掉：
    上一句就此蒸发，再也没人回答（"问了几句都没回应"的根因之一）。
    感官接入是"听见并记住"，与"要不要开口"是两件事，前者永不该被打断。
    """
    t = session.turn_task
    if t and not t.done():
        t.cancel()
        with contextlib.suppress(BaseException):
            await t
    session.turn_task = None
    a = session.ack_task
    if a and not a.done():
        a.cancel()
        with contextlib.suppress(BaseException):
            await a
    session.ack_task = None


# 摄像头/相册上传的画面帧：内存暂存（语音轮消费即删），容量约束只留最近8帧
_FRAMES: dict[str, bytes] = {}


def _store_frame(data: bytes) -> str:
    fid = uuid.uuid4().hex[:12]
    if len(_FRAMES) >= 8:
        _FRAMES.pop(next(iter(_FRAMES)), None)
    _FRAMES[fid] = data
    return fid


# ────────────────────────── 一轮对话全闭环 ──────────────────────────

async def _life_progress(owner: int) -> str:
    """人生脉络进度（注入 guide prompt）：按用户年龄裁剪阶段 + 真实种子聚合出每段已找回条数

    进度永远来自 alaya_seeds 聚合（习气网络同一批底层种子），不落任何冗余状态表。
    """
    rows = await db.pool().query(
        "SELECT u.age, u.birthday FROM xd_ais a JOIN xd_users u ON u.id=a.user_id WHERE a.id=%s",
        [int(owner)])
    age = life_stages.age_from(rows[0]) if rows else life_stages.DEFAULT_AGE
    cnt_rows = await db.pool().query(
        "SELECT life_stage, COUNT(*) c FROM alaya_seeds "
        "WHERE owner_ai_id=%s AND state IN ('confirmed','temp') AND life_stage IS NOT NULL "
        "GROUP BY life_stage", [int(owner)])
    counts = {str(r["life_stage"]): int(r["c"] or 0) for r in cnt_rows}
    focus = life_stages.pick_focus(age, counts)
    return life_stages.progress_text(age, counts, focus)


async def _skill_prompt(owner: int) -> str:
    """对话技能种子 → prompt 注入（仿真人对话的「通识能力」）

    技能种子是本体记忆通识库里的 skill 类种子（等待说完/话中回应/停顿控制/情绪标签/
    情绪音色），冷启动由 manas.ensure_traits 播种，此后由熏习演化。它们是「我会怎么说话」
    的能力，注入 system prompt 让 LLM 知道自己在这些技能上的习得——不是硬编码行为规则，
    而是把本体已习得的对话技能作为"我"的一部分交给 LLM 去灵活运用。
    """
    rows = await db.pool().query(
        "SELECT seed_text, raw_text FROM alaya_seeds "
        "WHERE owner_ai_id=%s AND seed_type='skill' AND state IN ('confirmed','temp') "
        "ORDER BY weight DESC, id DESC", [int(owner)])
    if not rows:
        return ""
    lines = []
    for r in rows:
        name = str(r.get("seed_text") or "").strip()
        desc = str(r.get("raw_text") or "").strip()
        if name and desc:
            lines.append(f"- {name}：{desc}")
        elif name:
            lines.append(f"- {name}")
    if not lines:
        return ""
    return ("【对话技能（你已习得的说话本领，灵活运用，不必生硬照做）】\n"
            + "\n".join(lines))


async def _self_model_for(ai_id: int, peer_ai_id: int) -> dict:
    """某个 AI 对本轮对话对象的完整我识：本体我 + 关系我 + 当下我。

    三层我识的存储位置不同（见 selfmodel）：
      core（本体我）存于 (ai_id, 0) —— 原则不动，与具体对话对象无关
      relation / mood 存于 (ai_id, peer_ai_id) —— 按对话对象分层
    只查 (ai_id, peer) 会漏掉 core（AI 便不记得自己是谁），必须合并。
    """
    sm = await selfmodel.get_self_model(int(ai_id), 0)
    if peer_ai_id:
        for layer in ("relation", "mood"):
            v = await selfmodel.get_layer(int(ai_id), int(peer_ai_id), layer)
            if v:
                sm[layer] = v
    return sm


def _fmt_gap(seconds: float) -> str:
    """一段时长 → 人话（连续时间维度的表述）"""
    s = max(0, int(seconds))
    if s < 60:
        return f"{s} 秒"
    if s < 3600:
        return f"{s // 60} 分钟"
    if s < 86400:
        h, m = divmod(s // 60, 60)
        return f"{h} 小时 {m} 分" if m else f"{h} 小时"
    return f"{s // 86400} 天"


async def _time_note(owner: int, peer: int, session: "Session | None" = None) -> str:
    """连续时间维度：把「现在几点、这轮聊了多久、距上次说话多久」**如实**告诉 AI。

    只报事实、不作判断——该顺着刚才的话头接着说，还是像重新见面一样起个新话头，
    由主念与眼下这句话的关联自己决定。系统**不设"间隔多久算新会话"之类的硬规则**
    （念头永续生发；用时间阈值去切分会话，等于用代码替心智做决定）。
    """
    now = datetime.now()
    weekday = "一二三四五六日"[now.weekday()]
    parts = [f"现在是 {now.strftime('%Y-%m-%d %H:%M')}（周{weekday}）。"]
    if session and session.started_at:
        age = (now - session.started_at).total_seconds()
        if age >= 120:
            parts.append(f"这轮对话你们已经聊了 {_fmt_gap(age)}。")
    rows = await db.pool().query(
        "SELECT TIMESTAMPDIFF('SECOND', created_at, NOW()) AS gap FROM alaya_dialogues "
        "WHERE owner_ai_id=%s AND peer_ai_id=%s ORDER BY id DESC LIMIT 1",
        [int(owner), int(peer)])
    if rows and rows[0].get("gap") is not None:
        parts.append(f"距你上一次和对方说话，已经过了 {_fmt_gap(max(0, int(rows[0]['gap'])))}。")
    return "【时间感】" + "".join(parts)


async def _build_context(owner: int, peer: int, user_text: str, speaker_name: str = "",
                         mode: str = "persona", ai_name: str = "", owner_name: str = "",
                         login_name: str = "", guest_note: str = "", time_note: str = "",
                         manas_core=None, autonomous: bool = False,
                         materials: list | None = None) -> tuple[str, dict]:
    """末那识主念集群（为主）+ 向量反射（补充）→ system prompt

    对话前提的唯一来源是末那识心智涌现：主念集群 = 「我此刻脑子里浮着的是什么」，
    由习气网络权重演化 + 全网竞争真实涌现，不是检索拼装出来的 prompt。
    向量反射只作补充——补上"与当下这句话直接相关、但尚未冲进主念集群"的记忆。

    记忆主体（谁的我识/记忆被注入）：
      guide  ：owner = 用户本人的 AI 本体——TA的记忆以第三人称档案注入给出场的如意看
      persona：peer  = 出场的对方 AI 本体——以本人格第一人称读自己的记忆"""
    mem_subject = owner if mode == "guide" else (peer or owner)
    mem_other = peer if mode == "guide" else owner

    dominant = list(getattr(manas_core, "dominant", None) or [])
    # 第六意识·综合分别：把末那识的主念集群收摄成「当下这一念」（受·想·思 + 伴念）
    conception = manas_core.conceive() if manas_core else {}
    rows = await retrieval.fetch_seed_rows([d["seed_id"] for d in dominant])

    # 向量补充：只补主念集群之外的直接相关记忆（主念已有的不重复）
    # ⚠️ 自主开口（没人说话）不做三路召回：近因 + 高权重常青会塞进来一批旧事，
    #    模型就顺着那批旧事开口（实测十几分钟里反复提同一件），而不是说"当下这一念"。
    #    自主轮的所缘只有主念集群（conception），这一念是什么就说什么。
    in_cluster = {int(d["seed_id"]) for d in dominant}
    if autonomous:
        recall = {"seeds": [], "concepts": [], "seed_ids": []}
    else:
        recall = await retrieval.recall(mem_subject, mem_other, user_text)
        hits = [s for s in (recall.get("seeds") or []) if int(s.get("seed_id") or 0) not in in_cluster]
        if hits:
            top = max(0.5, float(hits[0].get("score") or 0.5))
            await xixi.reinforce_on_recall(
                [{"seed_id": s["seed_id"], "confidence": min(1.0, float(s.get("score") or 0.5) / top)}
                 for s in hits if s.get("seed_id")])
        recall = {"seeds": hits, "concepts": recall.get("concepts") or [],
                  "seed_ids": [s["seed_id"] for s in hits]}

    # 屏幕上说话的那个本体 = 中枢归属者（末那识跟随出场者）：
    #   guide  ：跟如意对话，中枢与「我识」都是如意的——如意是一个 AI 本体，不是旁观的向导；
    #            她的习气网络就是此刻在想事的那个大脑，控制台切到她看到的就是这份实况。
    #   persona：亲友的 AI 本体本人格出场，中枢与「我」自然是它自己的。
    # 记忆双写（如意的记忆档案 + 用户本体的记忆档案）由 _memory_pipeline 的视角映射负责，与此无关。
    speaker_ai = int(peer or owner)
    speaker_self = await _self_model_for(speaker_ai, owner)
    my_self = selfmodel.build_self_prompt(speaker_self)

    if mode == "guide":
        ai_name = ai_name or "系统AI"
        owner_name = owner_name or "TA的本体"
        parts = [f"你的名字是「{ai_name}」。", GUIDE_PROMPT.format(owner_name=owner_name)]
        # 说话方式是她的硬规则（由她自己的 style 种子凝聚而成），单独抬出来强调，
        # 否则混在长上下文里会被"访谈要做细"的惯性淹没。
        # 她的完整我识不在此单独注入——统一与「当下这一念」合流（见下方 my_self）。
        habit = str((speaker_self.get("core") or {}).get("speech_habit") or "").strip()
        if habit:
            parts.append(f"【说话方式·必须严格遵守】{habit}")
        # 人生脉络：按年龄裁剪 + 真实种子聚合，指明本轮优先聊哪一段
        lp = await _life_progress(owner)
        if lp:
            parts.append(lp)
        if login_name:
            who = f"对方已登录账号，昵称「{login_name}」，这是可靠的初始线索；请自然地与TA确认与延伸了解更多。"
            parts.append(who)
    else:
        # persona：出场的就是对方 AI 本体本人，以第一人称说话，读自己的记忆
        parts = [BASE_PROMPT.format(ai_name=ai_name or "本体")]
        if guest_note:
            parts.append(guest_note)

    if speaker_name:
        if mode == "guide":
            parts.append(f"声纹识别确认：现在说话的就是「{speaker_name}」本人。")
        else:
            parts.append(f"声纹识别确认：现在说话的是「{speaker_name}」。")
    elif guest_note:
        # 亲友本人格出场：面前的人身份已明确（见上方 guest_note），不再走"未知对方"分支
        pass
    else:
        parts.append("你现在还不知道对方是谁（声纹未匹配）。没听清或没理解对方说什么时，自然地追问确认，绝不装懂；对方自我介绍后要记住称呼。")
    # 连续时间维度：现在几点、距上次说话过了多久、这轮聊了多久（没有它，AI 会顺着几小时前的旧话题说）
    if time_note:
        parts.append(time_note)
    # 对话技能种子（等待说完/话中回应/停顿/情绪音色…）注入：谁在说话，就是谁的习得
    skill_p = await _skill_prompt(speaker_ai)
    if skill_p:
        parts.append(skill_p)
    # ── 记忆块：当下这一念（第六意识收摄，「我」在其中）为主；guide 另附 TA 的记忆档案 ──
    cp = retrieval.build_conception_prompt(conception, rows, subject_label="我", self_prompt=my_self)
    if cp:
        parts.append(cp)
    elif my_self:
        # 当下无主念（网络静默流转）：此刻没有"在想的事"，但「我」仍须在场——
        # 我识先于记忆（见 selfmodel.build_self_prompt）。否则本体连自己是谁都会短暂失忆。
        parts.append("【我】" + my_self)
    mp = retrieval.build_memory_prompt(recall, autonomous=autonomous)
    if mp:
        if mode == "guide":
            # 召回的种子归属用户本体（mem_subject=owner），是"TA的往事"——第三人称档案，绝不认领
            parts.append(f"【{owner_name}的记忆档案】以下是TA本人的往事，供你了解TA——"
                         f"绝不能把这些说成你自己的经历：\n" + mp)
        else:
            parts.append(mp)
    if materials:
        # 资料 ≠ 对方说的话：这是对方通过对话窗口发给你的文件/照片内容，
        # 你已经看过、知道内容。必须与"对方口述"区分开，否则会把文件里的字当成 TA 的心里话。
        blocks = []
        for m in materials:
            kind = "照片（画面描述）" if str(m.get("kind") or "") == "image" else "文件"
            blocks.append(f"——《{m.get('name') or kind}》{kind}：\n{m.get('text') or ''}")
        parts.append("【对方发给你的资料】以下内容是对方通过对话窗口发给你的文件/照片，"
                     "你正在看它，已经知道里面写了什么。这是资料原文，不是对方口述的话，"
                     "也不是你自己说过的话——回答时可以直接引用里面的内容：\n"
                     + "\n".join(blocks)[:MATERIAL_PROMPT_LIMIT])
    if autonomous:
        # 自主开口（内生念头）：没人问，是你脑子里浮起了「当下这一念」。
        # 必须点明这一轮的性质，否则 LLM 会顺着上一轮自己那句话往下编（"AI 回答自己的回答"），
        # 说出来的就不是此刻想起的事，而是拿旧话头硬续——那就成了"随口说"。
        parts.append("【此刻是你自己想说话】没有人问你任何问题——是「当下这一念」自己浮上来的。"
                     "顺着这一念说你此刻想的、想说的；"
                     "不要复述或续写你上一轮说过的话，也不要假装对方刚问了你什么。")
        # 主动搭话的分寸（实测 2026-10-24 实测斥为"蠢"：她拿旧记忆当话题，
        # 用"你刚说那事儿，我记着呢"这种回应式开头凭空编出一段并不存在的对话）。
        parts.append("【主动搭话的分寸·必须遵守】"
                     "① 开头不许出现「你刚说」「你刚才说」「你刚才提到」「你之前说的那事」——"
                     "刚才根本没人跟你说话，那是回应式的话，用在本轮就是凭空编造；"
                     "② 不许复盘、追问、质问TA的往事（不许问「是不是那段时间太难了」这种揣度）；"
                     "③ 只说此刻你脑子里浮现的、轻松自然的一件小事，或对当下的关心；"
                     "④ 拿不准的记忆宁可不提——记忆里没有把握的事，就别说成TA经历过。")
        # 在场感知：人要主动搭话，先得知道"身边有没有人、是谁"。
        # 这些信息一直在会话里（声纹认出的说话人 / 登录昵称 / 档案称呼），
        # 过去没交给模型，它就只能对着空气说话——把记忆里想起的事说成"你刚说的"。
        who = str(speaker_name or login_name or owner_name or "").strip()
        if who:
            parts.append(
                f"【此刻你身边】「{who}」就在旁边，但TA此刻没有跟你说话——"
                f"你是自己待着的时候想起了这件事。开口要像人：先轻轻叫一声或确认TA在"
                f"（比如「哎，{who}」「{who}，你在吗」），再把你想起来的事说出去；"
                f"绝不能说成是TA刚告诉你的，也不要问「你刚才说什么」。")
        else:
            parts.append(
                "【此刻你身边】没有确认到具体的人。没人在旁边时，这一念只是你自己在想——"
                "要么不出声，要么就是自言自语；不要假装在跟谁对话。")
    return "\n\n".join(parts), recall


async def _load_history(owner: int, peer: int, limit: int = 0) -> list:
    n = int(limit or HISTORY_TURNS)
    rows = await db.pool().query(
        "SELECT role, content FROM alaya_dialogues WHERE owner_ai_id=%s AND peer_ai_id=%s "
        "ORDER BY id DESC LIMIT %s", [int(owner), int(peer), n])
    return [{"role": r["role"], "content": r["content"]} for r in reversed(rows)]


async def _save_dialogue(owner: int, peer: int, role: str, content: str, episode_id: str = ""):
    await db.pool().insert(
        "INSERT INTO alaya_dialogues (owner_ai_id, peer_ai_id, role, content, episode_id) "
        "VALUES (%s,%s,%s,%s,%s)",
        [int(owner), int(peer), role, content, episode_id or None])


async def _write_view(owner: int, peer: int, spec: list[dict], episode: str,
                      sense: str = "voice") -> list[dict]:
    """把一份视角规格落库（种子 + 向量 + 共现 + 关联熏习）"""
    written: list = []
    seed_ids: list = []
    star_pids: set = set()          # 忆点星光：本轮讲到的人 → 与他相关的事自动成星
    for s in spec:
        pid = await extract.resolve_person(owner, s.get("person_name") or "")
        tone_id = await extract.match_emotion_tone(owner, s.get("emotion_word") or "")
        # 写入仲裁·同一件事识别：先看库里有没有这颗记忆的「孪生记录」。
        # 同一件事被反复讲述时，不再让每条记录都独占一份分量，而是建 corr 边连成一族——
        # 任一条被想起，全族一起浮起（修「说过却想不起来 + 重复乱记」的根）。
        twin_id, twin_sim = await retrieval.nearest_seed(s["seed_text"])
        conf = float(s.get("confidence", 0.8))
        # 理解先于写入的代码级保证：抽取器自己都没把握的信息不落库。
        # extract 的判级口径里「寒暄级信息再确定也不超过 0.7」，低于 0.6 即「听不清/转述存疑」——
        # 这类东西一旦入库就成了"乱记"的源头，宁缺勿滥（实测 2026-10-04 反馈的根因之一）。
        if conf < EXTRACT_MIN_CONF:
            continue
        is_repeat = bool(twin_id) and twin_sim >= retrieval.MERGE_SIM
        if is_repeat:
            conf = min(conf, 0.7)
        seed_id = await seeds_mod.write_seed(owner, {
            "seed_type": s["seed_type"],
            "seed_text": s["seed_text"],
            "raw_text": s.get("raw_text") or "",
            "subject": s.get("subject") or "self",
            "person_id": pid,
            "seed_date": s.get("seed_date"),
            "time_fuzzy": s.get("time_fuzzy"),
            "emotion_tone_id": tone_id,
            "emotion_word": s.get("emotion_word"),
            "scene": s.get("scene"),
            "life_stage": s.get("life_stage"),
            "confidence": conf,
            "episode_id": episode,
            "sense_id": sense,
        }, peer_ai_id=peer)
        await retrieval.ensure_seed_vector(seed_id, s["seed_text"])
        if is_repeat:
            await seeds_mod.link_repeat(owner, seed_id, twin_id, twin_sim)
        seed_ids.append(seed_id)
        if pid:
            star_pids.add(int(pid))
        written.append({"seed_id": seed_id, "seed_type": s["seed_type"],
                        "seed_text": s["seed_text"], "subject": s.get("subject")})
    if seed_ids:
        await xixi.strengthen_cooccurrence(owner, seed_ids)
        await xixi.associate_seeds(owner, seed_ids)
    # 忆点星光：一颗星＝一个事件束。本轮讲到谁，就在谁的名下留一颗待亲友来点的星。
    for pid in star_pids:
        await seeds_mod.ensure_star(owner, pid, episode)
    return written


async def _write_dialogue_memory(owner: int, peer: int, dialogue: str, *,
                                 chat_name: str = "", speaker_name: str = "",
                                 sense: str = "voice", episode: str = "") -> dict:
    """一段对话 → 双视角长期记忆种子（「结生」与「忆点星光补星」共用的写入主体）。

    owner：记忆归属A——当前登录用户本人的 AI 本体（数字分身，ai_me_id）
    peer ：记忆归属B——出场的对方 AI 本体（系统AI如意 / 亲友的 AI 本体）
    chat_name   ：出场本体的称呼（事实里 ai 主体的称呼）
    speaker_name：用户本人的称呼（事实里 speaker 主体的称呼）
    episode     ：事件束锚。缺省自生成；忆点星光补星时由调用方传入该星的 episode。

    输入永远是**一段对话原文**（不是一轮的两段文本）——这是结生的关键：
    一次看全，"同一件事只出一条完整事实"这条铁律才真正落得下去。逐轮抽取时模型
    只看到一轮，同一件事被讲三遍就抽出三条互不相识的碎片，回忆时拼不出全貌。

    事件事实只抽一次（视角无关的中性陈述），再由确定性视角映射生成两份种子：
      A（owner，真人本体）：「我」的经历 → subject=self，第一人称
      B（peer，出场本体）：「我」的经历 → subject=self，第一人称（同事件归属相反）
    绝不允许把同一段文本原样写进两份——那是"同一条记忆复制两份"，不是双视角。
    """
    written: list = []
    episode = episode or uuid.uuid4().hex[:16]
    try:
        facts, retracted = await dualview.extract_facts_from_dialogue(
            dialogue, llm_chat_fn, ai_name=chat_name, speaker_name=speaker_name)
        if not facts:
            return {"written": [], "episode": episode, "new_persons": [], "dual": 0,
                    "retracted": retracted}

        speaker_label = speaker_name or "对方"
        chat_label = chat_name or "对方"

        # 视角A：owner = 用户本人的 AI 本体
        spec_a = dualview.spec_from_facts(facts, view="speaker",
                                          my_name=speaker_label, peer_name=chat_label)
        written = await _write_view(owner, peer, spec_a, episode, sense)

        # 视角B：peer = 出场 AI 本体
        dual = 0
        if peer > 0:
            spec_b = dualview.spec_from_facts(facts, view="ai",
                                              my_name=chat_label, peer_name=speaker_label)
            if spec_b:
                await _write_view(peer, owner, spec_b, episode, sense)
                dual = len(spec_b)
                # AI 的见解/领悟沉淀成它的价值取向——个性由它自己的认知养成，不靠 prompt 写死
                if any(s["seed_type"] in dualview.SELF_ONLY_KINDS for s in spec_b):
                    try:
                        await selfmodel.reform_values(peer, llm_chat_fn)
                    except Exception as e:
                        log.error("[memory] 价值取向重塑失败（不影响记忆写入）: %s", e)

        if written:
            retrieval.save_vectors()
            await organize.organize_after_turn(owner, [w["seed_id"] for w in written], llm_chat_fn)
        # 这段对话里获知的人物称呼（声纹注册用，数据通路传递，不做文本正则解析）
        new_persons = [s.get("person_name") for s in spec_a
                       if s.get("person_name") and s.get("person_name") != "对方"]
        return {"written": written, "episode": episode, "new_persons": new_persons,
                "dual": dual, "retracted": retracted}
    except asyncio.CancelledError:
        raise
    except Exception as e:
        log.error("[memory] 记忆写入失败: %s", e, exc_info=True)
        return {"written": written, "error": str(e)}


async def _retract_memories(ai_ids: list, texts: list, speaker_name: str = "") -> int:
    """用户明确否认/纠正过的内容 → 库里同义的旧种子作废（state=invalid）。

    被否认的事就是没发生过；它若留在记忆里，AI 会继续当真、继续主动提起
    （实测事故：一句 ASR 噪声被延伸成"服装厂"，用户当场否认，之后仍被反复提起）。
    这不是兜底——否认是用户给出的确定性事实，假记忆必须从检索与热态网络里消失。

    ⚠️ 双视角记忆库里同一件事有两份：「我曾在宏达机械厂工作三年」（用户本体）
    和「老陈曾在…」（如意本体）。被否认的事实是第三人称，直接和第一人称那份算
    相似度会掉到 0.52（实测），所以要把称呼换成"我"再算一遍——同一件事的两种人称
    都必须命中。
    """
    ids = sorted({int(x) for x in (ai_ids or []) if int(x or 0) > 0})
    texts = [str(t).strip() for t in (texts or []) if str(t).strip()]
    if not ids or not texts:
        return 0
    variants: list[str] = []
    for t in texts:
        variants.append(t)
        if speaker_name and speaker_name in t:
            variants.append(t.replace(speaker_name, "我"))
    vecs = await retrieval._encode([v[:200] for v in variants])
    cand: dict[int, float] = {}
    for q in vecs:
        for sid, sim in retrieval._cosine_topk(q, 8):
            if float(sim) >= RETRACT_SIM:
                cand[int(sid)] = max(cand.get(int(sid), 0.0), float(sim))
    if not cand:
        log.info("[retract] 未对上库里任何记忆：%s", texts)
        return 0
    ph = ",".join(["%s"] * len(ids))
    n = 0
    done_owners: set = set()
    for sid, sim in sorted(cand.items(), key=lambda kv: -kv[1]):
        rows = await db.pool().query(
            f"SELECT state, owner_ai_id, episode_id FROM alaya_seeds "
            f"WHERE id=%s AND owner_ai_id IN ({ph})", [int(sid), *ids])
        if not rows or str(rows[0].get("state")) != "confirmed":
            continue
        owner_of = int(rows[0]["owner_ai_id"])
        if owner_of in done_owners:
            continue          # 每个本体只按最相似的那一条定位，防误伤
        # 同一事件束（双视角的两份共享 episode）一起作废：
        # 被否认的事实在用户本体那边是"我…"，在如意那边是"老陈…"，
        # 两份必须同时消失，否则另一份还会被拿来提。
        episode = str(rows[0].get("episode_id") or "")
        targets = [int(sid)]
        if episode:
            twins = await db.pool().query(
                f"SELECT id FROM alaya_seeds WHERE episode_id=%s AND state='confirmed' "
                f"AND owner_ai_id IN ({ph})", [episode, *ids])
            targets = sorted({int(r["id"]) for r in twins} | {int(sid)})
        for tid in targets:
            await db.pool().exec(
                "UPDATE alaya_seeds SET state='invalid', weight=0 WHERE id=%s AND state='confirmed'",
                [int(tid)])
            retrieval.drop_vector(int(tid))
            for ai in ids:
                core = manas.get_core(ai)
                if core:
                    core.forget(int(tid))
            n += 1
            log.info("[retract] 作废被否认的记忆 #%s（sim=%.2f）", tid, sim)
        done_owners.add(owner_of)
    return n


async def _evolve_self(owner: int, peer: int, dialogue: str, peer_distill: bool = True) -> dict:
    """我识演化：这段对话带来的心境/关系变化——每轮都跑（心境是"当下"，不该等到结生）。

    双本体各自演化：A（用户本体）经历这段对话，B（出场本体）也经历这段对话。
    """
    dist = await selfmodel.distill(owner, peer, dialogue, llm_chat_fn, trigger_by="post_chat")
    if peer > 0 and peer_distill:
        await selfmodel.distill(peer, owner, dialogue, llm_chat_fn, trigger_by="post_chat")
    return {"distilled": dist}


async def _memory_pipeline(owner: int, peer: int, user_text: str, reply_text: str,
                           chat_name: str = "", speaker_name: str = "",
                           sense: str = "voice", peer_distill: bool = True,
                           episode: str = "") -> dict:
    """单轮对话 → 长期记忆种子 + 我识演化（忆点星光补星等"即时写入"场景用）。

    主对话**不走这里**：主对话的长期记忆由「结生」（_condense）按现量窗口整合写入，
    否则同一件事会被逐轮刻成碎片。
    """
    dialogue = f"【{speaker_name or '对方'}】{user_text or ''}\n【{chat_name or 'AI'}】{reply_text or ''}"
    out = await _write_dialogue_memory(owner, peer, dialogue, chat_name=chat_name,
                                       speaker_name=speaker_name, sense=sense, episode=episode)
    if out.get("retracted"):
        await _retract_memories([owner, peer], out["retracted"], speaker_name=speaker_name)
    out.update(await _evolve_self(owner, peer, dialogue, peer_distill=peer_distill))
    return out


async def _condense(session, reason: str, ws: WebSocket | None = None) -> dict:
    """结生：把现量窗口整理成长期记忆（对应人脑睡眠期的系统巩固）

    窗口里是"这一段真实讲过的话"（对方说的 + 本体自己说的），一次看全 →
    同一件事只出一条完整事实。写入走双视角映射（两个本体各存一份），
    写完后整批清理这一段临时记忆——临时记忆不留在阿赖耶识里。

    触发见 working.should_condense（静默 / 满载 / 突显门控），会话结束再补一次。
    """
    owner, peer = int(session.owner_ai_id), int(session.peer_ai_id)
    speaker_ai = _manas_ai(session)          # 窗口归属：念头涌现的那个本体
    items = await working.window(speaker_ai, owner)
    if not items:
        return {"written": [], "condensed": 0}
    # 满载时只消费最早的一段，后写入的留给下一轮——一次喂给整理意识的原文不宜过长
    batch = items[:working.CONDENSE_MAX_ITEMS]
    log.info("[working] 结生开始（%s）本体#%s × #%s：窗口 %s 条 ids=%s",
             reason, speaker_ai, owner, len(batch), [int(i["id"]) for i in batch][:12])
    dialogue = working.dialogue_text(batch)
    # 结生要调 LLM（抽取 + 两次我识演化），必须等用户静默后才跑——否则它会跟"下一句话"
    # 抢同一个 LLM 端点（对话首字 2s → 十几秒）。等不到就下次静默再结生，原料不会丢。
    async with _bg_llm_slot("结生"):
        out = await _write_dialogue_memory(owner, peer, dialogue,
                                           chat_name=session.ai_name,
                                           speaker_name=session.speaker_name or session.peer_label,
                                           sense=working.sense_of(batch))
        if out.get("error"):
            # 结生没成功就绝不清理原料——临时记忆是库里唯一的那一份，
            # 清掉就真丢了。下次触发（静默/满载）自然接着结生。
            log.warning("[working] 结生未完成（%s），临时记忆保留待下次：%s", reason, out["error"])
            return {"written": [], "condensed": 0}
        # 用户在这段里明确否认/纠正过的内容：把库里同义的旧种子作废（被否认的事没发生过）
        if out.get("retracted"):
            await _retract_memories([owner, peer], out["retracted"],
                                    speaker_name=session.speaker_name or session.peer_label)
        # 声纹锚点绝不在这里回填：结生时"新出现的人名"是模型从对话里猜的，拿它去绑定
        # 声纹就等于把某段声音硬塞给猜出来的人（本人锚点被别人的声音覆盖，记忆全乱）。
        # 本人声纹只在 _identify_speaker 确认认主时登记（见 _bootstrap_own_voice）。
        await _evolve_self(owner, peer, dialogue, peer_distill=(session.mode != "guide"))
    await working.drop_first(speaker_ai, owner, len(batch))
    written = out.get("written") or []
    log.info("[working] 结生（%s）本体#%s × #%s：%s 条临时记忆 → %s 颗长期种子",
             reason, speaker_ai, owner, len(batch), len(written))
    if ws is not None and written:
        with contextlib.suppress(Exception):
            await ws.send_json({"type": "memory", "written": written, "episode": out.get("episode")})
    return {"written": written, "condensed": len(batch)}


async def _run_turn(session: Session, ws: WebSocket, user_text: str, scene: str = "",
                    frame_id: str = "", proactive: bool = False):
    """一次自主表达（可被打断取消）：主念涌现驱动 → 检索 → 流式生成+流式合成 → 写回

    非回合制：本函数不再由"收到输入"直接触发，而是由末那识中枢在
    主念在全网竞争中胜出、统治力突破 DOM_MANIFEST（真正想说）时调度。user_text 是入网的感官内容，
    作为当下语境；是否开口已由中枢自主决定。
    """
    owner, peer = session.owner_ai_id, session.peer_ai_id
    t0 = time.monotonic()
    # 注意：感官内容（用户原话）已由 _sense_from_text/_sense_from_audio 即时下发 stt，
    # 这里绝不再发——否则中枢自主调度本函数时会把用户原话重复冒泡，且与 AI 回复交错，
    # 造成"AI 的话显示到用户侧"的错乱。

    frame = _FRAMES.pop(frame_id, b"") if frame_id else b""
    if frame and not scene:
        scene = await vl_describe(frame)     # 摄像头帧 → 场景描述（TA亲眼所见，先看后答）

    now_text = f"{user_text}\n（此刻我看到的画面：{scene}）" if scene else user_text
    # 这一轮是"她自己想说话"还是"回应对方"：由调用方（表达调度）给出，也兼容 user_text 为空。
    autonomous = bool(proactive or not now_text)
    if session.started_at is None:
        session.started_at = datetime.now()      # 本轮会话起点（时间感的另一半：这轮聊了多久）
    time_note = await _time_note(owner, peer, session)
    t_note = time.monotonic()
    system_prompt, recall = await _build_context(
        owner, peer, now_text, speaker_name=session.speaker_name,
        mode=session.mode, ai_name=session.ai_name, owner_name=session.owner_name,
        login_name=str((session.user or {}).get("nickname") or "") if session.mode == "guide" else "",
        guest_note=session.guest_note, time_note=time_note,
        manas_core=session.manas_core, autonomous=autonomous,
        materials=session.materials)
    t_ctx = time.monotonic()
    # 自主开口（内生念头）不给任何历史：没人问，话头只能来自"此刻想起的那一念"。
    # 喂历史（哪怕只有自己那两条）都会把模型拽去续上一轮的旧话头——实测它于是用
    # "你刚说那事儿，我记着呢"这种回应式开头，凭空编出一段并不存在的对话。
    history = [] if autonomous else await _load_history(owner, peer)
    # 本轮事件束 id：user/assistant 两条流水共享，双视角抽取据此成对取原始对话
    turn_episode = uuid.uuid4().hex[:16]
    if scene:
        # 画面独立成行入流水：此后每一轮的历史里，TA都"记得"自己看过这一幕
        await _save_dialogue(owner, peer, "user", f"【画面】{scene}", turn_episode)
    # 有外部输入（用户说了话/发了画面）才记入对话流水的 user 侧。
    # 纯念头自主涌现（user_text 为空）不是"用户说了句空话"——若把空串写进流水，
    # 下一轮 history 里就会出现一条空 user 消息，LLM 收到「自己上一句 + 空 user」
    # 便会顺着自己的话继续编，表现为"AI 回答自己的回答"。
    if user_text or scene:
        await _save_dialogue(owner, peer, "user", user_text, turn_episode)

    messages = [{"role": "system", "content": system_prompt}, *history]
    # 同样地，只有存在真实外部输入时才追加 user 消息；自主涌现轮不追加，
    # 让 LLM 基于历史与 system 自主延续念头，而不是对着空气说话。
    if now_text:
        messages.append({"role": "user", "content": now_text})
    t_prep = time.monotonic()

    full = ""
    buf = ""
    sent_visible = ""      # 已下发的可见文本（剥离标记后），保证流式字符不重不漏
    tts_q: asyncio.Queue = asyncio.Queue()
    # 情绪音色·随情换声：本轮 TTS 音色。LLM 在回复开头声明 [[情绪:xx]]，
    # 服务端据此查情绪音色锚点（me_c_tones）切换；无映射/无标记则用常色。
    turn_voice = session.voice
    t_first = [0.0, 0.0]   # 首个 LLM 字 / 首块音频（时延诊断用）

    async def tts_pump():
        """边生成边合成：LLM 一切出整句立即合成下发，不等全文"""
        while True:
            sent = await tts_q.get()
            if sent is None:
                return
            async for pcm in tts_stream(sent, turn_voice):
                if not t_first[1]:
                    t_first[1] = time.monotonic()
                await ws.send_bytes(pcm)

    def _visible(full_text: str) -> str:
        """流式可见文本：同时剥离寻忆标记与情绪标记（两者都是给系统的指令）"""
        return _visible_emotion_stream(_visible_stream(full_text))

    pump = asyncio.create_task(tts_pump())
    # 对话这一刻起，后台 LLM（自我演化/结生/资料抽取）必须让路——同一个 LLM 端点，
    # 它们在场就会把首字从 2 秒拖到十几秒（见 _bg_llm_slot 的实测注释）。
    _preempt_bg_llm()
    if autonomous:
        # 让界面知道这一句不是"回应"——是她自己主动搭话，气泡上标「主动」。
        await ws.send_json({"type": "proactive"})
    try:
        async for delta in llm_stream(messages, temperature=0.8, max_tokens=1024):
            if not t_first[0]:
                t_first[0] = time.monotonic()
            full += delta
            # 情绪标记一旦闭合即解析并切换本轮音色——必须早于第一句送入 TTS，
            # 否则整轮仍用常色，技能形同虚设。标记在开头，第一句切分前必已闭合。
            if turn_voice == session.voice:
                emo = _parse_emotion_marker(full)
                if emo:
                    tone_voice = extract.emotion_voice(emo)
                    if tone_voice:
                        turn_voice = tone_voice
            visible = _visible(full)
            if len(visible) > len(sent_visible):
                chunk = visible[len(sent_visible):]
                sent_visible = visible
                buf += chunk
                await ws.send_json({"type": "reply_delta", "text": chunk})
                while True:
                    sent, buf = cut_sentence(buf)
                    if not sent:
                        break
                    tts_q.put_nowait(sent)
        if buf.strip():
            tts_q.put_nowait(buf)
        await tts_q.put(None)
        await pump
    except BaseException:
        with contextlib.suppress(Exception):
            pump.cancel()   # 打断/异常：合成立即停，孤儿任务不残留
            await pump
        raise
    finally:
        if full.strip():
            # 打断时已生成部分也入对话流水，保上下文连续
            with contextlib.suppress(Exception):
                await _save_dialogue(owner, peer, "assistant",
                                     _strip_emotion_marker(_strip_recall_marker(full)),
                                     turn_episode)

    # 系统AI如意建议转发寻忆链接：剥离标记后，单独下发 ask_recall 触发前端底窗表单
    recall_ask = _parse_recall_marker(full)
    full = _strip_emotion_marker(_strip_recall_marker(full))
    # 分段时延：一段一行，谁慢一眼就能看出来（用户问"到底慢在哪"就看这里）
    log.info("[latency] 回复 %.2fs：时间感 %.2fs · 记忆/主念 %.2fs · 历史+落库 %.2fs · "
             "LLM首字 %s · 首块音频 %s · LLM全文 %.2fs",
             time.monotonic() - t0, t_note - t0, t_ctx - t_note, t_prep - t_ctx,
             f"{t_first[0] - t0:.2f}s" if t_first[0] else "-",
             f"{t_first[1] - t0:.2f}s" if t_first[1] else "-",
             time.monotonic() - t_prep)
    await ws.send_json({"type": "reply_done", "text": full, "recalled": len(recall.get("seeds", []))})
    if recall_ask:
        await ws.send_json({"type": "ask_recall", **recall_ask})
    # 纯念头自主涌现（无外部输入）不写记忆：这一轮没有"对方说的新事"。有外部输入时：
    #   ① 本体自己回的这一句也进现量窗口（只留档，不回头当"境"刺激网络）；
    #   ② 长期记忆种子与「我识演化」都不在这里做——交给「结生」把这一段一次看全后整合。
    #      逐轮就刻会刻出互不相识的碎片（"父亲守在奶奶身边"讲三遍变三条），
    #      我识也会被单轮的情绪带得来回抖动。
    if now_text:
        core = session.manas_core
        if core and full.strip():
            with contextlib.suppress(Exception):
                await core.inject_sensory(f"{session.ai_name}：{full[:400]}", sense_id="text",
                                          raw_text=full, to_network=False, peer_ai_id=owner,
                                          evidence=working.MARK_SELF)
    # 这一轮到此结束：turn_task 立刻完成，下一句"想说"马上能被调度。
    # ⚠️ 如意的自我养成（末那识凝聚 + 风格重凝）必须放到 turn_task 之外：
    #    它们每次都要调 LLM（实测一轮 30~60 秒）。留在 turn_task 里，用户紧接着说
    #    第二句时 _push 的 turn_task.done() 门控一直不放行——实测第二句等了 59 秒
    #    才开口，而服务端自己的分段时延只有 4 秒。这是"还是很慢"的真正大头。
    await ws.send_json({"type": "turn_done"})
    _spawn_guide_evolve(session, owner, user_text, full)


# 每本体同一时间只跑一份自我演化：连着说十句就排十份 LLM 演化任务只会互相拖死，
# 演化本来就是"事后回想"，跳过的那一轮会由下一次对话一并带上。
_evolve_tasks: dict[int, asyncio.Task] = {}


def _spawn_guide_evolve(session: Session, owner: int, user_text: str, full: str):
    """如意的自我养成（relation/mood 演化 + 说话风格沉淀）——后台跑，不占用表达通道"""
    ai_id = int(session.chat_ai_id or 0)
    if session.mode != "guide" or ai_id <= 0:
        return
    running = _evolve_tasks.get(ai_id)
    if running is not None and not running.done():
        log.info("[guide] 上一轮自我演化仍在跑（%s），本轮跳过", ai_id)
        return
    _evolve_tasks[ai_id] = _spawn_bg(_guide_evolve(ai_id, owner, user_text, full))


async def _guide_evolve(ai_id: int, owner: int, user_text: str, full: str):
    try:
        async with _bg_llm_slot("如意自我演化"):
            gd = await selfmodel.distill(ai_id, owner, f"【对方】{user_text}\n【我】{full}",
                                         llm_chat_fn, trigger_by="guide_turn", style_evolve=True)
            if gd.get("style_seed"):
                if await guide.harvest_style_seed(ai_id, gd["style_seed"]):
                    await guide.consolidate_style(ai_id, llm_chat_fn, trigger_by="guide_turn")
    except asyncio.CancelledError:
        log.info("[guide] 自我演化让路（对话优先），本轮跳过")
    except Exception as e:
        log.error("[guide] 如意凝聚失败: %s", e)
    finally:
        _evolve_tasks.pop(int(ai_id), None)


_INVALID_PERSON_NAMES = {"对方", "（未识别）", "未识别", "未知", "某人", "路人", "用户", "本人", "自己"}


def _valid_person_name(name) -> bool:
    """可注册/可入记忆的称呼：非空、不在无效名单、长度合理（防抽取LLM把占位词当人名）"""
    s = str(name or "").strip()
    return bool(s) and s not in _INVALID_PERSON_NAMES and len(s) <= 20


# 系统AI如意建议转发寻忆链接的标记：[[RECALL:人物称呼|关系|人生阶段]]（关系/阶段可空）
_RECALL_RE = re.compile(
    r"\[\[\s*RECALL\s*[:：]\s*([^\]|]*?)\s*"
    r"(?:\|\s*([^\]|]*?)\s*)?"
    r"(?:\|\s*([^\]]*?)\s*)?"
    r"\]\]")


def _strip_recall_marker(text: str) -> str:
    """剥离寻忆标记：这行是给系统的指令，绝不能让用户看到/听到"""
    return _RECALL_RE.sub("", str(text or "")).rstrip()


# 流式期间尚未闭合的标记：`[[` / `[[R` / `[[RECALL` / `[[RECALL:老王|同学|sch` …
_RECALL_PREFIX = "[[recall"


def _visible_stream(full: str) -> str:
    """流式可见文本：剥离完整标记，并**扣住**末尾可能长成标记的半截

    LLM 是一边生成一边下发的，标记在生成过程中一定会先出现半截（如 `[[RECALL:老王|同`）。
    此时正则匹配不到、会被当成正文发给用户，等标记补齐后又从可见文本里消失——
    用户那边就永久留下了一段残缺的 `[[RECALL:…`。所以在流式阶段要把这类尾巴先按住。
    """
    v = _strip_recall_marker(full)
    i = v.rfind("[[")
    if i < 0:
        return v
    tail = v[i:]
    if "]]" in tail:                     # 已闭合——完整标记已被上面剥掉，这里是普通文本
        return v
    probe = tail.replace(" ", "").lower()
    if _RECALL_PREFIX.startswith(probe) or probe.startswith(_RECALL_PREFIX) or probe.startswith("[[："):
        return v[:i]                     # 正在敲 / 已开始敲但没闭合：按住不发
    return v


def _parse_recall_marker(text: str) -> dict:
    """解析寻忆标记 → {person, relation, stage}；人物称呼无效则返回空（不弹窗、不编造）"""
    m = _RECALL_RE.search(str(text or ""))
    if not m:
        return {}
    person = str(m.group(1) or "").strip()[:20]
    if not _valid_person_name(person):
        return {}
    stage = str(m.group(3) or "").strip().lower()[:24]
    if stage not in life_stages.STAGE_KEYS:
        stage = ""
    return {"person": person, "relation": str(m.group(2) or "").strip()[:12], "stage": stage}


# 情绪标记：[[情绪:喜悦]]（情绪音色·随情换声技能的外显）
# 技能种子教会 LLM「情绪变化时切换音色」，LLM 在回复里用这个标记声明本轮情绪，
# 服务端据此查 me_c_tones 情绪音色锚点切换 TTS 音色。无映射则用常色（不编造）。
_EMOTION_RE = re.compile(r"\[\[\s*情绪\s*[:：]\s*([^\]]{1,12}?)\s*\]\]")
_EMOTION_PREFIX = "[[情绪"


def _strip_emotion_marker(text: str) -> str:
    """剥离情绪标记：这行是给系统的指令，绝不能让用户看到/听到"""
    return _EMOTION_RE.sub("", str(text or "")).rstrip()


def _visible_emotion_stream(full: str) -> str:
    """流式可见文本：剥离完整情绪标记，并扣住末尾可能长成标记的半截"""
    v = _strip_emotion_marker(full)
    i = v.rfind("[[")
    if i < 0:
        return v
    tail = v[i:]
    if "]]" in tail:
        return v
    probe = tail.replace(" ", "")
    if _EMOTION_PREFIX.startswith(probe) or probe.startswith(_EMOTION_PREFIX):
        return v[:i]
    return v


def _parse_emotion_marker(text: str) -> str:
    """解析情绪标记 → 情绪词（如"喜悦"）；无标记返回空"""
    m = _EMOTION_RE.search(str(text or ""))
    if not m:
        return ""
    return str(m.group(1) or "").strip()[:12]


# ────────────────────────── WS 语音会话网关 ──────────────────────────

async def ws_converse(ws: WebSocket):
    await ws.accept()
    session = Session()
    log.info("[ws] 连接建立")
    try:
        while True:
            msg = await ws.receive()
            if msg["type"] == "websocket.disconnect":
                break
            if msg.get("bytes"):
                session.audio_buf += msg["bytes"]
                continue
            text = msg.get("text")
            if not text:
                continue
            try:
                data = json.loads(text)
            except ValueError:
                continue
            mtype = data.get("type")

            if mtype == "init":
                # 强制登录：init 必须携带 user_id/token；对话AI归属本人或公共AI
                u = await auth.ws_user(data.get("user_id"), data.get("token"))
                if not u:
                    await ws.send_json({"type": "error", "message": "请先登录", "code": "auth"})
                    await ws.close()
                    continue
                session.user = u
                inv_tok = str(data.get("invite_token") or "").strip()
                if inv_tok:
                    # 忆点星光·guest：亲友凭已接受邀请与对方本体对话（owner=本体，记忆写本体）
                    inv = await db.pool().query("SELECT * FROM alaya_galaxy_invites WHERE token=%s", [inv_tok])
                    if not inv or inv[0]["status"] != "accepted" or int(inv[0]["invitee_user_id"] or 0) != int(u["id"]):
                        await ws.send_json({"type": "error", "message": "邀请无效或不属于当前用户", "code": "auth"})
                        await ws.close()
                        continue
                    ai_rows = await db.pool().query(
                        "SELECT id, name, voice_id FROM xd_ais WHERE id=%s", [int(inv[0]["inviter_ai_id"])])
                    if not ai_rows:
                        await ws.send_json({"type": "error", "message": "邀请指向的AI本体不存在"})
                        continue
                    # 出场本体B = 邀请人的 AI 本体（"忆点星光"）；记忆归属A = 被邀请人本人的 AI 本体
                    session.chat_ai_id = int(ai_rows[0]["id"])
                    session.mode = "persona"
                    session.peer_ai_id = int(ai_rows[0]["id"])
                    session.owner_ai_id = await auth._ensure_user_ai(int(u["id"]))
                    session.owner_name = session.ai_name = ai_rows[0]["name"]
                    if data.get("voice") is None and ai_rows[0].get("voice_id"):
                        session.voice = str(ai_rows[0]["voice_id"] or "8")
                    addr = str(inv[0]["address_term"] or "").strip()
                    rel = str(inv[0]["relation_type"] or "好友").strip()
                    session.speaker_name = addr or str(u.get("nickname") or "")
                    session.peer_label = session.speaker_name
                    # 本次寻忆挂靠的人生阶段（邀请时选定）：让本体知道"要补的是哪一段经历"
                    st_key = str(inv[0].get("life_stage") or "").strip()
                    st_name = life_stages.stage_name(st_key)
                    st_note = (f"这次TA来，是想帮你补全「{st_name}」那一段经历——"
                               f"围绕那个时期的共同往事展开。" if st_name else "")
                    session.guest_note = (
                        f"对方是「{addr or u.get('nickname')}」（你是TA的{rel}），TA通过「忆点星光」来找你——"
                        "你只握着这段往事的自己那一半，另一半得靠对方讲给你听。"
                        + st_note +
                        "多听多问细节（时间/地点/当时情形），真诚表达记忆被一点点补全的感受；绝不编造记忆里没有的细节。")
                    await retrieval.load_vectors_once()
                    await ws.send_json({"type": "ready", "ai_name": session.ai_name, "voice": session.voice,
                                        "mode": "persona", "owner_name": session.owner_name})
                    continue
                session.chat_ai_id = int(data.get("owner_ai_id") or 0)
                if session.chat_ai_id <= 0:
                    await ws.send_json({"type": "error", "message": "缺少 owner_ai_id"})
                    continue
                try:
                    await auth.require_owner_ai(u, session.chat_ai_id)
                except Exception:
                    await ws.send_json({"type": "error", "message": "无权访问该AI本体", "code": "auth"})
                    await ws.close()
                    continue
                await retrieval.load_vectors_once()
                rows = await db.pool().query(
                    "SELECT id, name, voice_id, user_id FROM xd_ais WHERE id=%s", [session.chat_ai_id])
                if not rows:
                    await ws.send_json({"type": "error", "message": f"AI #{session.chat_ai_id} 不存在"})
                    continue
                session.ai_name = rows[0]["name"]
                if data.get("voice") is None and rows[0].get("voice_id"):
                    session.voice = str(rows[0]["voice_id"] or "8")
                # 模式判定（判据铁律）：
                #   对话永远发生在「两个 AI 本体」之间：
                #     A = 当前登录用户本人的 AI 本体（数字分身，xd_users.ai_me_id）——记忆归属A
                #     B = 出场的对方 AI 本体（如意 / 亲友的AI本体，即本次选择的 chat_ai_id）——记忆归属B
                #   用户绝不能与自己的 AI 本体对话，只能与「其它」AI 本体聊。
                #   绝不能拿 xd_ais.user_id == xd_users.id 当本体判据——那只是「归属人」，
                #   用户名下的角色/亲友分身(user_id 相同)统统会被误判成本体。
                #   公共AI（user_id=0，如意）= guide：如意以向导角色出场，注入的档案是用户本人的记忆
                #   亲友的 AI 本体           = persona：该本体本人格出场，读自己的记忆
                chat_owner = int(rows[0].get("user_id") or 0)
                my_ai_id = await _user_main_ai(u["id"], fallback=0) or await auth._ensure_user_ai(int(u["id"]))
                if not my_ai_id:
                    await ws.send_json({"type": "error", "message": "当前账号的AI本体不可用", "code": "no_self_ai"})
                    await ws.close()
                    continue
                if int(session.chat_ai_id) == int(my_ai_id):
                    await ws.send_json({"type": "error",
                                        "message": "不能与自己的AI本体对话，请选择其它AI本体",
                                        "code": "self_ai"})
                    await ws.close()
                    continue
                session.owner_ai_id = int(my_ai_id)              # 记忆归属A：用户本人的 AI 本体
                session.peer_ai_id = int(session.chat_ai_id)     # 记忆归属B：出场的对方 AI 本体
                session.speaker_name = str(u.get("nickname") or "")
                session.peer_label = session.speaker_name        # 用户本人的称呼（视角兜底）
                if chat_owner == 0:
                    # 公共系统AI（如意）：向导角色出场，记忆档案 = 用户本人的本体
                    session.mode = "guide"
                    own_rows = await db.pool().query(
                        "SELECT name FROM xd_ais WHERE id=%s", [session.owner_ai_id])
                    session.owner_name = own_rows[0]["name"] if own_rows else "TA的本体"
                else:
                    # 亲友的 AI 本体：该本体本人格出场（忆点星光），读自己的记忆
                    session.mode = "persona"
                    session.owner_name = session.ai_name
                    rel = await db.pool().query(
                        "SELECT relation_type, address_term, life_stage FROM alaya_galaxy_invites "
                        "WHERE status='accepted' AND ((inviter_user_id=%s AND invitee_user_id=%s) "
                        "OR (inviter_user_id=%s AND invitee_user_id=%s)) LIMIT 1",
                        [chat_owner, int(u["id"]), int(u["id"]), chat_owner])
                    rel_t = str((rel[0].get("relation_type") if rel else "") or "亲友").strip()
                    addr = str((rel[0].get("address_term") if rel else "") or "").strip()
                    who = addr or str(u.get("nickname") or "这位朋友")
                    st_name = life_stages.stage_name(
                        str((rel[0].get("life_stage") if rel else "") or ""))
                    st_note = (f"这次TA想帮你补全的是「{st_name}」那一段经历，请围绕那个时期展开。"
                               if st_name else "")
                    session.guest_note = (
                        f"面前的人是「{who}」——TA是你的{rel_t}，你们之间有过共同的往事。"
                        "你丢失了大量记忆，正在努力想起自己是谁；TA是来帮你恢复记忆的。"
                        + st_note +
                        "请围绕「请TA帮你回忆你们共同经历的事」展开对话：主动问起你们一起经历过的场景、"
                        "时间、地点、当时的人和情形，把TA讲的每一件往事当作找回自己的线索。"
                        "如果已有相关记忆浮现，可以自然聊起来、逐渐熟络，共同建立新的记忆。"
                        "但绝不能被TA带着跑题——始终紧扣「找回你自己的记忆」这条主线，"
                        "不聊与你们共同经历无关的闲事，绝不编造记忆里没有的细节。")
                await ws.send_json({"type": "ready", "ai_name": session.ai_name, "voice": session.voice,
                                    "mode": session.mode, "owner_name": session.owner_name})
                # 末那识心智涌现中枢：本体上线 → 热态载入 + 人格特质种子就位 + 启动状态推送
                await _manas_attach(session, ws)

            elif mtype == "start":
                # 用户开始说话：进入聆听态——AI 不抢话，等 TA 说完（见 _push 的抢话门控）。
                # ⚠️ 绝不清空 audio_buf：用户中途停顿后继续说，客户端会重发 start，
                # 旧实现无条件 audio_buf=b"" 会把前半句整段丢掉，只剩最后半句
                # （"前面说了半天全丢了"的根因）。缓冲按聆听段连续累积，只在 stop 提交时清空。
                session.listening = True
                session.listen_started_at = time.monotonic()
                session.silence_ms = 0.0
                session.ack_cursor = len(session.audio_buf)   # 话中回应从本段新音频起算
                # 用户一开口，后台 LLM（演化/结生/资料抽取）立刻让路，别跟对话抢算力
                _preempt_bg_llm()
                # 打断 AI 正在说的话：用户开口即打断（生成/写回随任务取消终止）
                await _cancel_turn(session)

            elif mtype == "interim":
                # 话中回应：对方还在说、中途停顿换气（客户端判到短停顿，非说完）。
                # 取出本段新增音频应一声「嗯/对/继续」表示在听——是否应声由主念自主涌现决定。
                # ⚠️ 不清空 audio_buf：用户还在继续说，这段音频仍是本回合的一部分，
                # 最终 stop 时要连同后续一起识别（否则又变成"只录到最后半句"）。
                seg = session.audio_buf[session.ack_cursor:]
                session.ack_cursor = len(session.audio_buf)
                if seg:
                    session.ack_task = asyncio.create_task(_ack_from_audio(session, ws, seg))

            elif mtype == "stop":
                # 用户说完（客户端 VAD 判到静音）：退出聆听态，提交本段语音识别。
                session.listening = False
                session.silence_ms = float(data.get("silence_ms") or 0.0)
                audio = session.audio_buf
                session.audio_buf = b""
                session.ack_cursor = 0
                frame_id = str(data.get("frame_id") or "")
                # 非回合制：识别后只把感官内容注入习气网络，是否开口由主念自主涌现决定
                session.sense_task = asyncio.create_task(_sense_from_audio(session, ws, audio, frame_id))

            elif mtype == "chat":
                text_in = str(data.get("text") or "").strip()
                scene = str(data.get("scene") or "").strip()
                if not text_in and not scene:
                    continue
                if not text_in:
                    text_in = "（发来一张照片）"
                # 非回合制：文字/画面入网为感官种子，不强制应答
                session.sense_task = asyncio.create_task(_sense_from_text(session, ws, text_in, scene))

            elif mtype == "material":
                # 对话窗口发来的资料（文档正文 / 照片画面描述）：进上下文 + 进记忆，
                # 但绝不当成"用户说的话"（见 _sense_material）。
                kind = str(data.get("kind") or "text")
                name = str(data.get("name") or "")
                body = str(data.get("text") or "").strip()
                if not body:
                    await ws.send_json({"type": "notice", "text": "这份资料没解析出内容，换一份试试"})
                    continue
                session.sense_task = asyncio.create_task(
                    _sense_material(session, ws, kind, name, body))

            elif mtype == "interrupt":
                await _cancel_turn(session)

            elif mtype == "voice":
                # 控制台里切换了该本体的音色：本会话立刻改用新音色说话
                vid = str(data.get("voice") or "").strip()
                if vid:
                    session.voice = vid
                    await ws.send_json({"type": "voice_ok", "voice": vid})

            elif mtype == "manas":
                # 前端查询中枢实时状态（激活值快照 + 阈值层级 + 主念集群）
                core = manas.get_core(_manas_ai(session)) if _manas_ai(session) else None
                if core:
                    await ws.send_json({"type": "manas_state", **core.snapshot_payload()})
    except WebSocketDisconnect:
        pass
    finally:
        await _cancel_turn(session)
        if session.sense_task and not session.sense_task.done():
            session.sense_task.cancel()
            with contextlib.suppress(BaseException):
                await session.sense_task
        if session.manas_task:
            session.manas_task.cancel()
        # 会话下线：回收它留下的"想说"意愿——人走了，那句话不再是想回谁的依据，
        # 免得别的会话（或本会话重连后的新 sid）把"跟一个已经离开的人说话"排成表达。
        if session.manas_core:
            session.manas_core.release_session(session.sid)
        # 会话结束：把还没结生的现量窗口巩固掉——人也是聊完"睡一觉"完成巩固。
        # 临时记忆是落了库的，所以即便这里是异常断开也能结生，不会"一断连啥也没记住"。
        with contextlib.suppress(Exception):
            if session.condense_task:
                await session.condense_task
            await _condense(session, "会话结束")
        log.info("[ws] 连接断开")


def _manas_ai(session) -> int:
    """末那识归属本体 = 屏幕上说话的那个 AI 本体（出场者）。

    跟如意对话，中枢就是如意的——如意是一个 AI 本体，有自己的习气网络与我识，
    念头就在她的网络里涌现；跟亲友的 AI 本体对话，中枢就是那个本体的。
    这不是"向导旁白"：谁在说话，就该看到谁的大脑实况。
    记忆双写（同一件事各写一份到对方本体）由 _memory_pipeline 的视角映射负责，与中枢归属无关。
    """
    return int(session.peer_ai_id or session.owner_ai_id or 0)


async def _manas_attach(session, ws: WebSocket):
    """本体上线：载入热态习气网络（中枢自此永不停止迭代）+ 周期性推送激活值增量

    非回合制核心：中枢每 tick 自主迭代，一旦主念在全网竞争中胜出且统治力突破
    DOM_MANIFEST（真正想说），就由这里调度一次自主表达——开口与否、何时开口，
    全由念头自主涌现决定。
    """
    core = await manas.ensure_core(_manas_ai(session))
    session.manas_core = core
    await ws.send_json({"type": "manas_state", **core.snapshot_payload()})

    async def _push():
        last_tick = 0
        last_sent_tick = 0
        last_condense_tick = 0
        while True:
            # ⚠️ 与中枢同拍（每 tick 一次），不是每 5 tick——"想说"意愿的寿命由那颗感官种子
            # 决定（当下之境 9 秒），但取它的时机必须紧跟产生它的那一拍。过去每 5 拍（1 秒）
            # 才来看一次，意愿经常在两次轮询之间就被图层的抖动吃掉（"说了话没回应"的根因）。
            await asyncio.sleep(manas.TICK_MS / 1000)
            # 自主表达调度：主念突破阈值 → 取走"想说"意愿 → 以上下文开口
            # 顺序铁律：① 先判断 turn_task 空闲（ASR/声纹识别期间不抢），② 再取走意愿，
            # ③ 取到后立即 commit_expression 标记感官种子已消费。忙碌时保留意愿，
            # 等空闲后自然取走——否则"想说"意愿被丢弃，用户这句话就再也没人回答。
            # ④ 聆听门控：用户此刻正在说话（listening）绝不开口——等 TA 说完。
            #    仿真人对话的礼貌：人不会在对方话说到一半时插嘴。意愿保留在 pending，
            #    用户说完（stop）后下一拍自然取走，绝不丢。
            # ⑤ 归属门控：意愿只回给「接收了那句话」的会话（session.sid），
            #    页面挂着但没说话的会话取不走——避免回复跑到错的屏幕上。
            if not session.listening \
                    and (session.sense_task is None or session.sense_task.done()) \
                    and (session.turn_task is None or session.turn_task.done()):
                # ⚠️ take_expression 必须在"本会话确实空闲"之后调用：它会消费掉意愿，
                # 忙碌时取走就等于把这次"想说"丢了（感官种子已提交、却没人回答）。
                expr = core.take_expression(session.sid)
            else:
                expr = None
            if expr is not None:
                # 意愿确实被调度出去，此刻才标记感官种子已消费（见 manas.commit_expression）。
                # 若在 take_expression 内标记，则忙碌时被丢弃的意愿会让感官种子永久消费，
                # 用户这句话在冷却期内永不重试——"没回复"的根因。
                # 语境必须跟着意愿的**来源**走：
                #   speech       —— 是用户这句话顶起来的，语境 = 那句话；
                #   spontaneous  —— 她自己想起来的，语境必须为空。
                # 过去无条件取 session.pending_sense，残留的上一句用户原话会被当成
                # "用户现在说的话"，如意于是对着旧话回应、把记忆里的事说成用户刚说的。
                if expr.get("source") == "speech":
                    sense, scene = session.pending_sense, session.pending_scene
                    core.commit_expression(session.sid)
                else:
                    sense, scene = "", ""
                    core.commit_expression("")      # 自主轮不消费任何"用户说的话"
                # 感官内容已被主念采纳（或本轮是自主念头），即刻清空——
                # 否则下一轮表达会重复采纳这段旧文本
                session.pending_sense = ""
                session.pending_scene = ""
                session.turn_task = asyncio.create_task(
                    _run_turn(session, ws, sense, scene=scene,
                              proactive=(expr.get("source") != "speech")))
            # 结生（系统巩固）：静默 / 满载 / 突显门控三种情形，见 working.should_condense。
            # ① 每 50 拍（≈10 秒）查一次——窗口的边界是秒级的，不必每拍算；
            # ② 会话空闲时才动手（正在说话/正在生成时不打断）；
            # ③ 用任务异步跑：结生要调 LLM，绝不能卡住这里的中枢状态推送。
            if not session.listening \
                    and (session.condense_task is None or session.condense_task.done()) \
                    and core.tick - last_condense_tick >= 50:
                last_condense_tick = core.tick
                reason = await working.should_condense(_manas_ai(session), session.owner_ai_id)
                if reason:
                    session.condense_task = asyncio.create_task(_condense(session, reason, ws))
            if core.tick == last_tick:
                continue
            last_tick = core.tick
            # 中枢状态推送仍按 5 拍节流（每秒一次）——前端看的是趋势，不必每拍发
            if core.tick - last_sent_tick < 5:
                continue
            last_sent_tick = core.tick
            try:
                await ws.send_json({"type": "manas_state", **core.snapshot_payload()})
            except Exception:
                return

    session.manas_task = asyncio.create_task(_push())
    log.info("[ws] 末那识中枢已接入本体 #%s（%s 种子 / %s 边）",
             _manas_ai(session), len(core.nodes), len(core.edges))


async def _user_main_ai(user_id: int, fallback: int = 0) -> int:
    """用户记忆归属本体：xd_users.ai_me_id（登录即建的权威本体）> alaya_user_main > 名下第一个AI"""
    rows = await db.pool().query(
        "SELECT a.id FROM xd_users u JOIN xd_ais a ON a.id=u.ai_me_id "
        "WHERE u.id=%s AND a.is_alive=1", [int(user_id)])
    if rows:
        return int(rows[0]["id"])
    m = await db.pool().query(
        "SELECT m.main_ai_id FROM alaya_user_main m JOIN xd_ais a ON a.id=m.main_ai_id "
        "WHERE m.user_id=%s AND a.is_alive=1", [int(user_id)])
    if m:
        return int(m[0]["main_ai_id"])
    fr = await db.pool().query(
        "SELECT id FROM xd_ais WHERE user_id=%s AND is_alive=1 ORDER BY id LIMIT 1", [int(user_id)])
    return int(fr[0]["id"]) if fr else int(fallback or 0)


async def _sense_from_text(session: Session, ws: WebSocket, text_in: str, scene: str = ""):
    """感官接入（文字/画面）：入网为瞬时感官种子，作为当下语境暂存，
    是否开口完全交给末那识中枢自主涌现——本函数绝不触发应答。"""
    try:
        if not session.manas_core:
            return
        sense = "photo" if scene else "text"
        content = f"{text_in}\n（此刻我看到的画面：{scene}）" if scene else text_in
        who = session.speaker_name or session.peer_label or "对方"
        await session.manas_core.inject_sensory(
            f"{who}：{content}", sense_id=sense, raw_text=content,
            peer_ai_id=session.owner_ai_id, evidence=working.MARK_SPEAKER,
            session_key=session.sid)
        session.pending_sense = text_in
        session.pending_scene = scene
        await ws.send_json({"type": "stt", "text": text_in})
    except asyncio.CancelledError:
        raise
    except Exception as e:
        log.error("[sense] 文字接入失败: %s", e, exc_info=True)


async def _sense_material(session: Session, ws: WebSocket, kind: str, name: str, text: str):
    """对话窗口发来的资料（文档正文 / 照片画面描述）。

    资料 ≠ 对方说的话（实测 2026-10-24 定）：
      · 正文原样挂进会话上下文（session.materials），此后每一轮都读得到。
        旧实现把正文塞进 pending_sense 当"用户这句话"，一次打断（barge-in 取消 turn）
        就整篇蒸发——用户看到的就是"上传了文档，如意根本不知道里面写了什么"。
      · 入网只放一条摘要感官种子（让"想说"意愿起来），不把整篇正文当对方口述。
      · 全文走记忆导入同一条链路形成长期种子（与 /api/memory/import 同一份铁律）。
    """
    body = str(text or "").strip()
    if not body:
        return
    body = body[:MATERIAL_MAX_TEXT]
    nm = str(name or "").strip() or ("照片" if kind == "image" else "资料")
    session.materials = (session.materials + [{"kind": kind, "name": nm, "text": body}])[
        -MATERIAL_MAX_KEEP:]
    label = "一张照片" if kind == "image" else f"一份资料《{nm}》"
    if not session.manas_core:
        return
    who = session.speaker_name or session.peer_label or "对方"
    digest = f"{who}发来{label}：{body[:400]}"
    await session.manas_core.inject_sensory(
        digest, sense_id=("photo" if kind == "image" else "text"), raw_text=body,
        # 资料不入现量窗口（MARK_LISTEN 的角色就是"不写成一句对方说的话"）：
        # 它的长期记忆只走下面 _ingest_material 这一条链路，避免同一份资料被
        # 结生再抽取一遍（实测会写出一模一样的两颗种子）。
        peer_ai_id=session.owner_ai_id, evidence=working.MARK_LISTEN,
        session_key=session.sid)
    session.pending_sense = f"（我发来{label}，你看到了吗？）"
    session.pending_scene = ""
    # 全文进长期记忆：同一份写入链路（不另造类型、不做 RAG）。
    # ⚠️ 这里不进 _bg_tasks：资料是用户明确的动作，必须落库——只允许"等用户静默后再跑"，
    #    不允许被对话取消掉（对话来了它只是在门口等着，不占 LLM）。
    asyncio.create_task(_ingest_material(session.owner_ai_id, body))


async def _ingest_material(owner_ai_id: int, text: str):
    """资料全文 → 长期种子（与 /api/memory/import 完全同一条链路）"""
    try:
        async with _bg_llm_slot("资料进记忆"):
            specs, episode = await _import_specs(int(owner_ai_id), [], text)
            if not specs:
                return
            written = await _write_view(int(owner_ai_id), 0, specs, episode, "text")
            retrieval.save_vectors()
        log.info("[material] 本体#%s 资料正文 → %s 颗长期种子", owner_ai_id, len(written))
    except asyncio.CancelledError:
        raise
    except Exception as e:
        log.error("[material] 资料进记忆失败: %s", e, exc_info=True)


async def _sense_from_audio(session: Session, ws: WebSocket, audio: bytes, frame_id: str = ""):
    """感官接入（语音）：ASR 识别 → 入网为瞬时感官种子 + 暂存当下语境。

    无效语音（识别为空/垃圾/幻觉）不形成种子、不写习气网络、界面也不显示——
    像人一样，听到跟自己无关或没听清的话，就当没听见，不予理会。
    唯一例外：声音很大（响度超阈值）却语义不清时，形成一颗「疑问」种子——
    像人听到有人冲自己大声说了句什么却没听清，会心生疑惑。
    识别结果不直接触发应答；主念是否突破阈值由中枢自主决定。
    """
    try:
        if len(audio) < 6000:  # <0.19s 视为误触
            return
        t0 = time.monotonic()
        # 看图是最慢的一步（本地 VL 2~6 秒）。它跟"这段音频说了什么"完全独立，
        # 所以必须在 ASR 之前并行启动——原来是 ASR 完再串行看图，白白多等一次。
        frame = _FRAMES.pop(frame_id, b"") if frame_id else b""
        scene_task = asyncio.create_task(vl_describe(frame)) if frame else None
        user_text, spk_emb, rms = await asr_recognize(audio, speaker=True)
        t_asr = time.monotonic()
        if not user_text:
            if scene_task:
                scene_task.cancel()
                with contextlib.suppress(BaseException):
                    await scene_task
            # 无效语音：不入网、不下发（相当于没听见）。
            # 但若声音很大，形成疑问——「刚才有人大声说了句什么，我没听清」。
            if rms >= LOUD_RMS and session.manas_core:
                await session.manas_core.inject_sensory(
                    "刚才有人大声冲我说了句什么，我没听清——TA说了什么？",
                    sense_id="voice", gain=LOUD_GAIN,
                    peer_ai_id=session.owner_ai_id, evidence=working.MARK_LISTEN,
                    session_key=session.sid)
            return
        # 听声识人（铁律）：只有「当前登录用户本人」的声纹才算用户的声音。
        # 其他人的声音——不写记忆、不当成用户说的话、也不用用户的名义作答。
        is_me, who_name = await _identify_speaker(session, spk_emb)
        t_spk = time.monotonic()
        if not is_me:
            if scene_task:
                scene_task.cancel()
                with contextlib.suppress(BaseException):
                    await scene_task
            await ws.send_json({"type": "notice",
                                "text": (f"这句像是「{who_name}」说的，不是你——没有写进记忆"
                                         if who_name else
                                         "没辨认出是你本人的声音，这句没有记住（如意只认你本人的声纹）")})
            return
        if who_name:
            session.speaker_name = who_name
        await ws.send_json({"type": "stt", "text": user_text})
        if not session.manas_core:
            if scene_task:
                scene_task.cancel()
                with contextlib.suppress(BaseException):
                    await scene_task
            return
        # 画面帧：先看后想——视觉描述作为感官内容的一部分入网
        # （VL 任务已经跑了这么久，这里通常只剩很短的等待）
        scene = await scene_task if scene_task else ""
        content = f"{user_text}\n（此刻我看到的画面：{scene}）" if scene else user_text
        who = session.speaker_name or session.peer_label or "对方"
        await session.manas_core.inject_sensory(
            f"{who}：{content}", sense_id="voice", raw_text=content,
            peer_ai_id=session.owner_ai_id, evidence=working.MARK_SPEAKER,
            session_key=session.sid)
        session.pending_sense = user_text
        session.pending_scene = scene
        # 分段时延（用户问"到底慢在哪"时就靠这一行定位）：ASR / 识人 / 入网各占多少
        log.info("[latency] 语音接入 %.2fs 音频=%.1fs：ASR %.2fs · 识人 %.2fs · 入网(含语义关联) %.2fs",
                 time.monotonic() - t0, len(audio) / 32000.0,
                 t_asr - t0, t_spk - t_asr, time.monotonic() - t_spk)
    except asyncio.CancelledError:
        raise
    except Exception as e:
        log.error("[sense] 语音接入失败: %s", e, exc_info=True)


async def _ack_from_audio(session: Session, ws: WebSocket, audio: bytes):
    """话中回应：对方还在说、中途停顿换气时，应一声「嗯」「对」「继续」表示在听。

    仿真人对话的关键行为：人不会在对方长篇讲述时全程沉默，也不会抢话打断，
    而是在停顿的间隙用极简回应跟随。是否应声由末那识主念自主涌现决定——
    只有当主念集群里浮起「话中回应」技能种子（"我在听"这个念头胜出）才应一声；
    否则保持安静聆听。绝不写死"每段停顿必回嗯"。

    ⚠️ 与正式回复的区别：这里是**聆听态下的短回应**，不打断用户叙述、不进入
    LLM 长回复链路、不写对话流水（它不是一个回合，只是"在听"的表示）。
    """
    try:
        if not session.manas_core:
            return
        if len(audio) < 6000:      # <0.19s 视为误触
            return
        # 话中回应不需要声纹（只是"我在听"的应声），故按 speaker=False 口径解包 2 个返回值
        text, _rms = await asr_recognize(audio)
        if not text:
            return
        # 入网为感官种子：让"对方正在说什么"进入当下语境，参与主念竞争。
        # 不进现量窗口（MARK_LISTEN）：这只是聆听中途的片段，stop 时还会提交完整句。
        await session.manas_core.inject_sensory(
            str(text)[:500], sense_id="voice", raw_text=text,
            peer_ai_id=session.owner_ai_id, evidence=working.MARK_LISTEN,
            session_key=session.sid)
        # 主念自主决定：此刻浮起的是「话中回应」技能，才应一声
        if not session.manas_core.ack_in_dominant():
            return
        ack = await _pick_ack(session, text)
        if not ack:
            return
        async for pcm in tts_stream(ack, session.voice):
            await ws.send_bytes(pcm)
    except asyncio.CancelledError:
        raise
    except Exception as e:
        log.error("[ack] 话中回应失败: %s", e, exc_info=True)


async def _pick_ack(session: Session, said: str) -> str:
    """从「话中回应」技能里挑一句极简应声（嗯/对/继续/没错/是/哦？/啊？）

    用 LLM 按对方刚说的话挑最贴切的一句——不是随机，也不是写死。
    上下文极短、输出极短，走轻量通道。
    """
    sys_p = ("你在听对方讲述，对方中途停顿换气。你只需应一声表示在听、在跟随，"
             "不打断、不评论、不追问。从「嗯」「对」「继续」「没错」「是」「哦？」「啊？」"
             "里挑最贴切的一个，只输出这一个词，不要标点、不要别的话。")
    try:
        out = await llm_chat_fn([
            {"role": "system", "content": sys_p},
            {"role": "user", "content": f"对方刚说：{said}"},
        ], temperature=0.6, max_tokens=8)
        ack = str(out or "").strip().strip("。！？，、,.!? ")[:4]
        return ack if ack in ("嗯", "对", "继续", "没错", "是", "哦？", "啊？", "哦", "啊") else ""
    except Exception:
        return ""


# ── 听声识人（铁律·实测 2026-10-24）─────────────────────────────
#   对话里只认「当前登录用户本人」的声纹。其他人的声音一律不算用户的声音——
#   不写记忆、不当成"用户说的话"、也不以用户的名义作答，否则记忆全乱。
#   本人参考：xd_speaker 里绑定到该 user_id 的簇（全平台唯一权威）
#           + alaya_persons 里与本账号同名的锚点（本链路早期自动登记的簇）。
#   同域簇实测：同人 0.53~0.76，不同人 ≤0.40 —— 0.45 是可分界的下限。
SPEAKER_ME_FLOOR = 0.45      # 判为本人的相似度下限
SPEAKER_ME_MARGIN = 0.05     # 本人还须明显领先"最像的他人"，避免都像时误判
SPEAKER_EMB_DIM = 192        # CAM++ 声纹维度


def _unit(v):
    import numpy as np
    n = float(np.linalg.norm(v))
    return v / n if n > 0 else v


async def _speaker_refs(session: Session) -> tuple[list, list]:
    """(本人声纹参考列表, 他人参考[(称呼, 向量)])——他人只用于报出"那是谁"，不参与认主"""
    import numpy as np
    uid = int((session.user or {}).get("id") or 0)
    me_name = str((session.user or {}).get("nickname") or "").strip()
    me: list = []
    others: list = []
    nbytes = SPEAKER_EMB_DIM * 4
    if uid:
        rows = await db.pool().query(
            "SELECT embedding FROM xd_speaker WHERE user_id=%s AND embedding IS NOT NULL "
            "AND LENGTH(embedding)=%s", [uid, nbytes])
        me += [_unit(np.frombuffer(r["embedding"], dtype=np.float32)) for r in rows]
        rows = await db.pool().query(
            "SELECT nickname, embedding FROM xd_speaker WHERE user_id IS NOT NULL AND user_id<>%s "
            "AND embedding IS NOT NULL AND LENGTH(embedding)=%s", [uid, nbytes])
        for r in rows:
            rig = _unit(np.frombuffer(r["embedding"], dtype=np.float32))
            # 同一昵称的另一账号簇（同一个人开了两个号）也算本人参考，
            # 否则本人会被判成"旁人：自己"
            if me_name and str(r["nickname"] or "").strip() == me_name:
                me.append(rig)
            else:
                others.append((str(r["nickname"] or "旁人"), rig))
    ai_ids = sorted({int(x) for x in (session.owner_ai_id, session.peer_ai_id) if int(x or 0) > 0})
    if ai_ids:
        import base64 as b64
        ph = ",".join(["%s"] * len(ai_ids))
        rows = await db.pool().query(
            f"SELECT name, speaker_emb FROM alaya_persons "
            f"WHERE owner_ai_id IN ({ph}) AND speaker_emb IS NOT NULL", ai_ids)
        for r in rows:
            emb = np.frombuffer(b64.b64decode(r["speaker_emb"]), dtype=np.float32)
            if len(emb) != SPEAKER_EMB_DIM:
                continue
            name = str(r["name"] or "").strip()
            if me_name and name == me_name:
                me.append(_unit(emb))
            else:
                others.append((name or "旁人", _unit(emb)))
    return me, others


async def _identify_speaker(session: Session, spk_emb_b64: str | None) -> tuple[bool, str]:
    """声纹 → (是否当前用户本人, 称呼)。

    只有确认是「当前登录用户本人」才返回 True；其他人/未识别一律 False。
    绝不再把任何一个陌生声纹自动登记成本人（旧实现每说一句就覆盖一次本人锚点，
    于是谁说话都成了"老陈"——记忆错乱的根）。
    """
    if not spk_emb_b64:
        # 未部署声纹服务：单人使用场景，麦克风前的人就是当前登录用户本人。
        # 这是开放版的明确语义，不是比对失败后的兜底——配了声纹服务才会走下面的认主逻辑。
        if not settings.get("speaker.enabled"):
            return True, str((session.user or {}).get("nickname") or "").strip()
        log.warning("[speaker] 本轮没有声纹向量，无法确认是不是本人")
        return False, ""
    import numpy as np, base64 as b64
    cur = np.frombuffer(b64.b64decode(spk_emb_b64), dtype=np.float32)
    if len(cur) != SPEAKER_EMB_DIM:
        log.warning("[speaker] 声纹维度异常（%s），无法比对", len(cur))
        return False, ""
    cur = _unit(cur)
    me_refs, other_refs = await _speaker_refs(session)
    s_me = max((float(cur @ r) for r in me_refs), default=-1.0)
    s_other, other_name = -1.0, ""
    for name, ref in other_refs:
        s = float(cur @ ref)
        if s > s_other:
            s_other, other_name = s, name
    if me_refs and s_me >= SPEAKER_ME_FLOOR and s_me >= s_other + SPEAKER_ME_MARGIN:
        return True, str((session.user or {}).get("nickname") or "").strip()
    if not me_refs and s_other < SPEAKER_ME_FLOOR:
        # 本账号从未登记过本人声纹：第一位开口的人就是本人（认主），此后以它为准
        name = str((session.user or {}).get("nickname") or "").strip()
        if name:
            await _bootstrap_own_voice(session, spk_emb_b64, name)
            return True, name
    if s_other >= SPEAKER_ME_FLOOR:
        log.info("[speaker] 非本人：%s（本人最像 %.2f / 他人 %.2f）", other_name, s_me, s_other)
        return False, other_name
    log.info("[speaker] 未识别声纹（本人最像 %.2f / 他人最像 %.2f）——不作本人处理", s_me, s_other)
    return False, ""


async def _bootstrap_own_voice(session: Session, spk_emb_b64: str, name: str):
    """首次认主：本账号还没有任何本人声纹参考时，把第一段声音登记为本人的锚点"""
    for ai in sorted({int(x) for x in (session.owner_ai_id, session.peer_ai_id) if int(x or 0) > 0}):
        pid = await extract.resolve_person(ai, name)
        await db.pool().exec("UPDATE alaya_persons SET speaker_emb=%s WHERE id=%s",
                             [spk_emb_b64, int(pid)])
    session.speaker_name = name
    log.info("[speaker] 本人声纹首次登记：%s", name)


# ────────────────────────── 应用组装 ──────────────────────────

async def _decay_scheduler():
    """每日 04:00 半衰调度（越常用越稳固，久未唤醒按半衰期下沉）"""
    while True:
        now = datetime.now()
        target = now.replace(hour=4, minute=0, second=0, microsecond=0)
        if target <= now:
            target += timedelta(days=1)
        await asyncio.sleep((target - now).total_seconds())
        try:
            n = await xixi.decay_half_life()
            log.info("[scheduler] 每日半衰完成：%s 颗", n)
        except Exception as e:
            log.error("[scheduler] 半衰失败: %s", e)


async def _transient_gc():
    """临时记忆到期清理：结生是主路径（结生完就清了），这里清的是会话被意外中断、
    没能走到结生的残留。不清的话临时记忆会永久堆积在阿赖耶识里——实测过：
    过期的感官碎片长期滞留，还被"缺向量自愈"补进向量索引，当成往事召回注入 prompt。"""
    while True:
        await asyncio.sleep(60)
        try:
            n = await working.expired_gc()
            if n:
                log.info("[scheduler] 临时记忆到期清理 %s 条", n)
        except Exception as e:
            log.error("[scheduler] 临时记忆清理失败: %s", e)


async def _warm_vectors():
    """向量模型自检 + 索引回填（外部 API 就绪后调用；配置变更后也可再次触发）"""
    try:
        ok = await retrieval.probe()
    except Exception as e:
        log.error("[alaya] 向量模型不可用: %s（我的资料 → 模型设置）", e)
        return
    if not ok:
        return
    try:
        await retrieval.load_vectors_once()
    except Exception as e:
        log.error("[alaya] 向量索引回填失败: %s", e)


async def _warm_external():
    """外部服务的预热与自愈（后台执行）。

    服务可用性不依赖外部 API：接口先起来，向量自检/补算、如意我识凝聚在后台跑。
    配置还没填时这里会失败并在日志里说明，填好 key 后下一轮对话即可用。
    """
    await _warm_vectors()
    # 基线风格种子是如意的底色，不依赖任何外部服务，开机就种
    try:
        gid = await guide.guide_ai_id()
        await guide.ensure_guide_seeds(gid)
    except Exception as e:
        log.error("[alaya] 如意基线种子种入失败: %s", e)
    if not settings.is_configured("llm"):
        log.warning("[alaya] 未配置 LLM 接口，如意我识凝聚暂缓（我的资料 → 模型设置）")
        return
    try:
        gid = await guide.guide_ai_id()
        await guide.ensure_guide_ready(gid, llm_chat_fn)
    except Exception as e:
        log.error("[alaya] 如意初始化失败: %s", e)


@asynccontextmanager
async def lifespan(app: FastAPI):
    db.init_schema()
    # 公共本体（控制台里对所有登录用户开放的「如意」）
    try:
        await auth.ensure_public_ai()
    except Exception as e:
        log.error("[alaya] 公共本体初始化失败: %s", e)
    # 本地向量文件先读进来（毫秒级）；外部模型的探测/自愈放后台，不阻塞启动
    await retrieval.load_vectors()
    warm = asyncio.create_task(_warm_external())
    task = asyncio.create_task(_decay_scheduler())
    gc_task = asyncio.create_task(_transient_gc())
    # 末那识心智涌现中枢：永久后台自运行循环（毫秒级迭代，无外部输入也永不停止）
    manas_task = asyncio.create_task(manas.run_forever())
    log.info("[alaya] 引擎就绪（%s）", settings.get("server.port", 9890))
    yield
    for t in (warm, task, gc_task, manas_task):
        t.cancel()


app = FastAPI(title="Alaya Memory Engine", lifespan=lifespan)
app.include_router(auth.router)
app.include_router(profile.router)
app.include_router(amnesia.router)


@app.get("/api/status")
async def api_status(user: dict = Depends(auth.get_current_user)):
    seed_n = await db.pool().query("SELECT COUNT(*) AS n FROM alaya_seeds")
    return {
        "engine": "alaya-v2",
        "seeds": int(seed_n[0]["n"]) if seed_n else 0,
        "vectors": retrieval.vector_count(),
        "emb_ready": retrieval.model_ready(),
        "configured": {k: settings.is_configured(k)
                       for k in ("llm", "asr", "tts", "embedding")},
        "time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }


# ── 模型设置（控制台可改，写回 config.json，立即生效）──
_CONFIG_SECTIONS = ("llm", "vision", "asr", "tts", "embedding", "speaker")


def _strip_masked(patch: dict, cur: dict) -> dict:
    """把界面回显的掩码 key（****abcd）去掉，避免把星号写进配置"""
    out = {}
    for k, v in (patch or {}).items():
        if isinstance(v, dict):
            sub = _strip_masked(v, (cur or {}).get(k) or {})
            if sub:
                out[k] = sub
        elif k == "api_key" and isinstance(v, str) and v.startswith("*"):
            continue
        else:
            out[k] = v
    return out


@app.get("/api/config")
async def api_config_get(user: dict = Depends(auth.get_current_user)):
    """当前配置（api_key 只回显尾 4 位）+ 各段是否已配置"""
    cfg = settings.load()
    view = {k: cfg.get(k) for k in _CONFIG_SECTIONS}
    view["server"] = {"default_ai_name": cfg.get("server", {}).get("default_ai_name")}
    return {"config": settings.mask(view),
            "configured": {k: settings.is_configured(k) for k in _CONFIG_SECTIONS}}


@app.post("/api/config")
async def api_config_save(body: dict, user: dict = Depends(auth.get_current_user)):
    """保存模型设置（只接受白名单段；改完立即生效，无需重启）"""
    patch = _strip_masked({k: v for k, v in (body or {}).items() if k in _CONFIG_SECTIONS},
                          settings.load())
    srv = (body or {}).get("server")
    if isinstance(srv, dict) and "default_ai_name" in srv:
        patch["server"] = {"default_ai_name": str(srv["default_ai_name"] or "如意")}
    if not patch:
        return {"ok": True}
    settings.save(patch)
    log.info("[config] 模型设置已更新：%s", ",".join(patch.keys()))
    if "embedding" in patch or "llm" in patch:
        # 向量模型或 LLM 换了：后台重做自检 + 索引回填 + 如意我识凝聚
        asyncio.create_task(_warm_external())
    return {"ok": True, "configured": {k: settings.is_configured(k) for k in _CONFIG_SECTIONS}}


@app.get("/api/ais")
async def api_ais(user: dict = Depends(auth.get_current_user)):
    """可切换的对话角色：登录用户的AI本体 + 系统AI如意 + 亲友的AI本体。
    「AI本体」的唯一权威判据：xd_ais.id == xd_users.ai_me_id（登录即建）。
    绝不能用 xd_ais.user_id 当判据——那只是归属人标记，
    用户名下的角色、亲友分身会统统被误当成本体。
    公共AI（user_id=0）里只保留系统AI如意，其余系统AI不进入本应用的切换清单。"""
    uid = int(user["id"])
    # 自己的AI本体（唯一判据 ai_me_id）+ 系统AI如意（公共AI里唯一保留的一个）
    rows = await db.pool().query(
        "SELECT * FROM ("
        "  SELECT a.id, a.name, a.avatar, a.voice_id, a.user_id FROM xd_ais a "
        "  JOIN xd_users u ON u.ai_me_id=a.id "
        "  WHERE u.id=%s AND a.is_alive=1 "
        "  UNION "
        "  SELECT id, name, avatar, voice_id, user_id FROM xd_ais "
        "  WHERE is_alive=1 AND is_public=1"
        ") t ORDER BY (t.user_id=%s) DESC, t.id LIMIT 100", [uid, uid])
    # 亲友的AI本体：双向已接受邀请关系里对端用户的 ai_me_id（唯一判据）
    kin = await db.pool().query(
        "SELECT DISTINCT a.id, a.name, a.avatar, a.voice_id, a.user_id "
        "FROM alaya_galaxy_invites i "
        "JOIN xd_users ku ON ku.id = IF(i.inviter_user_id=%s, i.invitee_user_id, i.inviter_user_id) "
        "JOIN xd_ais a ON a.id = ku.ai_me_id "
        "WHERE i.status='accepted' AND (i.inviter_user_id=%s OR i.invitee_user_id=%s) "
        "AND a.is_alive=1 AND a.user_id>0 "
        "ORDER BY a.id LIMIT 100", [uid, uid, uid])
    seen = {int(r["id"]) for r in rows}
    for r in kin:
        if int(r["id"]) not in seen:
            rows.append(r)
            seen.add(int(r["id"]))
    # 权威标记由后端下发：is_my_ai = 该AI是否就是当前用户的AI本体（ai_me_id），
    # 前端不得自行用 user_id / is_user_voice 推断（会误判角色与亲友分身为本体）。
    my_ai_id = await _user_main_ai(uid, fallback=0)
    for r in rows:
        r["is_my_ai"] = 1 if int(r["id"]) == int(my_ai_id) and my_ai_id else 0
        r["is_system"] = 1 if int(r.get("user_id") or 0) == 0 else 0
    return {"ais": rows, "my_ai_id": my_ai_id}


@app.get("/api/user/main_ai")
async def api_get_main_ai(user: dict = Depends(auth.get_current_user)):
    """我的默认记忆本体（guide 模式下系统AI访谈的记忆归属）"""
    return {"main_ai_id": await _user_main_ai(user["id"], fallback=0)}


@app.post("/api/user/main_ai")
async def api_set_main_ai(body: dict, user: dict = Depends(auth.get_current_user)):
    main_id = int(body.get("main_ai_id") or 0)
    if main_id <= 0:
        return JSONResponse({"error": "main_ai_id 必填"}, status_code=400)
    rows = await db.pool().query(
        "SELECT id FROM xd_ais WHERE id=%s AND user_id=%s AND is_alive=1", [main_id, int(user["id"])])
    if not rows:
        return JSONResponse({"error": "只能选择自己名下的AI本体"}, status_code=403)
    await db.pool().exec(
        "INSERT INTO alaya_user_main (user_id, main_ai_id, updated_at) VALUES (%s,%s,NOW()) "
        "ON CONFLICT(user_id) DO UPDATE SET main_ai_id=excluded.main_ai_id, "
        "updated_at=excluded.updated_at",
        [int(user["id"]), main_id])
    return {"ok": True}


@app.get("/api/user")
async def api_user(owner_ai_id: int, user: dict = Depends(auth.get_current_user)):
    """当前AI本体归属的用户资料（控制台用户图标弹窗数据源，含登录手机号）"""
    await auth.require_owner_ai(user, owner_ai_id)
    rows = await db.pool().query(
        "SELECT u.nickname, u.phone, u.avatar, u.call_me, u.profession, u.gender, u.age, u.city, "
        "u.interests, u.birthday, u.zodiac "
        "FROM xd_ais a JOIN xd_users u ON u.id = a.user_id WHERE a.id=%s", [int(owner_ai_id)])
    return {"user": rows[0] if rows else None}


@app.post("/api/frame")
async def api_frame(file: UploadFile = File(...), user: dict = Depends(auth.get_current_user)):
    """摄像头截图上传：内存暂存返回 frame_id，语音轮 stop 消息携带，消费即删"""
    data = await file.read()
    if not data:
        return JSONResponse({"error": "空文件"}, status_code=400)
    return {"frame_id": _store_frame(data)}


def _parse_file_sync(lower: str, data: bytes) -> str:
    """文档文本提取（同步库，to_thread 中执行）"""
    import io
    if lower.endswith(".txt"):
        return data.decode("utf-8", errors="replace")
    if lower.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in reader.pages)
    if lower.endswith(".docx"):
        from docx import Document
        doc = Document(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    if lower.endswith(".pptx"):
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    t = shape.text_frame.text.strip()
                    if t:
                        parts.append(t)
        return "\n".join(parts)
    if lower.endswith((".doc", ".ppt")):
        raise ValueError("旧版二进制格式无法解析，请另存为 docx/pptx 后发送")
    raise ValueError("不支持的文件类型")


@app.post("/api/upload")
async def api_upload(file: UploadFile = File(...), user: dict = Depends(auth.get_current_user)):
    """多模态上传：图片→VLM场景描述；txt/pdf/docx/pptx→文本提取（前8000字）"""
    name = file.filename or "file"
    lower = name.lower()
    data = await file.read()
    if not data:
        return JSONResponse({"error": "空文件"}, status_code=400)
    try:
        if lower.endswith((".jpg", ".jpeg", ".png", ".webp", ".gif", ".bmp")):
            desc = await vl_describe(data)
            if not desc:
                return JSONResponse({"error": "画面识别失败"}, status_code=500)
            return {"kind": "image", "name": name, "desc": desc}
        text = await asyncio.to_thread(_parse_file_sync, lower, data)
        return {"kind": "text", "name": name, "text": text[:8000]}
    except ValueError as e:
        return JSONResponse({"error": str(e)}, status_code=400)


# ── 记忆导入：把"当下获取到的信息"形成记忆种子 ─────────────────────────
# 铁律（用户定，架构级，不得动摇）：
#   ① 绝不回退到 RAG —— 不做"外部知识库 + 检索注入"那一套。信息只有一种归宿：
#      阿赖耶识的记忆种子。
#   ② 绝不分类型对待 —— 不管导进来的是经历、资料、日记、知识还是照片，
#      一视同仁走**同一条写入链路**（写入仲裁 / 置信度 / 关联熏习 / 概念凝聚 / 向量 / 半衰），
#      不做任何特殊处理、不打任何特殊标记。
# 落地：它就是「用户对本体说了一段话，只是没有本体回复那一半」——
#   抽取用 dualview.extract_facts_from_dialogue（与对话共用同一份抽取铁律），
#   写入用 _write_view（与对话共用同一份写入链路）。零新增概念、零新增类型。
# 唯一差别：没有对话对端，peer_ai_id=0，只写这一份。
#
# 编码 = 种子规格本身（字段与 _write_view 的 spec 一一对应），不另造格式：
#   [{"text": "必填：这条记忆的内容",
#     "seed_type": "event",        // 可省，默认 event
#     "raw": "原话", "date": "2026-05-18", "time_fuzzy": "下午",
#     "person": "爸爸", "emotion": "高兴", "scene": "出行",
#     "life_stage": "midlife", "subject": "共有", "confidence": 1.0}]
IMPORT_MAX_ITEMS = 200
_IMPORT_KINDS = ("event", "value", "style", "thought",
                 "opinion", "insight", "imagination", "curiosity")
_IMPORT_SUBJECTS = ("self", "共有", "他人")


def _import_item_to_spec(it: dict) -> dict | None:
    """一条导入编码 → 种子规格（与 spec_from_facts 的输出同构，写入前无需转换）"""
    text = str(it.get("text") or "").strip()
    if not text:
        return None
    kind = str(it.get("seed_type") or "event").strip()
    subject = str(it.get("subject") or "self").strip()
    return {
        "seed_type": kind if kind in _IMPORT_KINDS else "event",
        "seed_text": text[:500],
        "raw_text": str(it.get("raw") or "").strip() or text,
        "subject": subject if subject in _IMPORT_SUBJECTS else "self",
        "person_name": str(it.get("person") or "").strip(),
        "seed_date": dualview.norm_date(it.get("date")),
        "time_fuzzy": str(it.get("time_fuzzy") or "").strip() or None,
        "emotion_word": str(it.get("emotion") or "").strip() or "平静",
        "scene": str(it.get("scene") or "").strip(),
        "life_stage": str(it.get("life_stage") or "").strip(),
        "confidence": max(0.3, min(1.0, float(it.get("confidence") or 1.0))),
    }


async def _owner_names(owner_ai_id: int) -> tuple[str, str]:
    """本体的归属人称呼 + 本体名

    导入时前者就是"讲这段内容的人"（等同对话里的 speaker），后者是"听的人"（等同 ai_name）。
    """
    rows = await db.pool().query(
        "SELECT u.nickname, a.name FROM xd_ais a JOIN xd_users u ON u.id=a.user_id WHERE a.id=%s",
        [int(owner_ai_id)])
    if not rows:
        return "我", "本体"
    speaker = str(rows[0].get("nickname") or rows[0].get("name") or "我").strip()
    return speaker, str(rows[0].get("name") or "本体").strip()


async def _import_specs(owner_ai_id: int, items: list, text: str) -> tuple[list, str]:
    """导入内容 → 种子规格 + 事件束锚

    结构化条目：直接用编码（用户给的就是准确的元信息，不必也不该让模型重写一遍）；
    一段文本/文件/照片描述：等同于"这个人对本体讲了这段内容"，按对话同一份铁律抽取。
    """
    episode = uuid.uuid4().hex[:16]
    if items:
        specs = [s for s in (_import_item_to_spec(i) for i in items if isinstance(i, dict)) if s]
        return specs, episode
    body = str(text or "").strip()
    if not body:
        return [], episode
    speaker, ai_name = await _owner_names(owner_ai_id)
    facts, _retracted = await dualview.extract_facts_from_dialogue(
        f"【{speaker}】{body[:6000]}", llm_chat_fn, ai_name=ai_name, speaker_name=speaker)
    # view="speaker"：这些是"这个人自己的事"，写进它自己的本体时主体就是「我」
    specs = dualview.spec_from_facts(facts, view="speaker", my_name=speaker, peer_name=ai_name)
    return specs, episode


@app.post("/api/memory/import")
async def api_memory_import(body: dict, user: dict = Depends(auth.get_current_user)):
    """记忆导入（JSON）：结构化条目直落，或一段文本按同一份铁律抽成种子。

    见上方铁律：不走 RAG、不分类型，与对话写入完全同一条链路。
    """
    owner_ai_id = int(body.get("owner_ai_id") or 0)
    if owner_ai_id <= 0:
        return JSONResponse({"error": "owner_ai_id 必填"}, status_code=400)
    await auth.require_owner_ai(user, owner_ai_id)
    items = body.get("items") or []
    if not isinstance(items, list):
        return JSONResponse({"error": "items 必须是数组"}, status_code=400)
    if len(items) > IMPORT_MAX_ITEMS:
        return JSONResponse({"error": f"单次最多导入 {IMPORT_MAX_ITEMS} 条，请分批"},
                            status_code=400)
    specs, episode = await _import_specs(owner_ai_id, items, str(body.get("text") or ""))
    if not specs:
        return {"written": [], "episode": episode}
    written = await _write_view(owner_ai_id, 0, specs, episode, str(body.get("sense") or "text"))
    retrieval.save_vectors()
    log.info("[import] 本体#%s 导入 %s 条编码 → 写入 %s 颗种子", owner_ai_id, len(specs), len(written))
    return {"written": written, "episode": episode}


def _photo_taken_at(data: bytes) -> str:
    """照片拍摄日期（EXIF DateTimeOriginal）——旅游照/生活照的核心元信息"""
    try:
        import io
        from PIL import Image
        exif = Image.open(io.BytesIO(data)).getexif()
        raw = str(exif.get(36867) or exif.get(306) or "").strip()  # DateTimeOriginal / DateTime
        if len(raw) >= 10 and raw[4] == ":" and raw[7] == ":":
            return raw[:10].replace(":", "-")
    except Exception:
        pass
    return ""


@app.post("/api/memory/import/file")
async def api_memory_import_file(owner_ai_id: int, file: UploadFile = File(...),
                                 user: dict = Depends(auth.get_current_user)):
    """记忆导入（文件/照片）：文档提文本、照片由 VLM 描述，二者都按同一份铁律形成种子"""
    await auth.require_owner_ai(user, owner_ai_id)
    name = (file.filename or "file").lower()
    data = await file.read()
    if not data:
        return JSONResponse({"error": "空文件"}, status_code=400)
    try:
        if name.endswith((".jpg", ".jpeg", ".png", ".webp", ".gif", ".bmp")):
            desc = await vl_describe(data)
            if not desc:
                return JSONResponse({"error": "画面识别失败"}, status_code=500)
            taken = _photo_taken_at(data)
            text, sense = (f"（拍摄于 {taken}）{desc}" if taken else desc), "photo"
        else:
            text, sense = await asyncio.to_thread(_parse_file_sync, name, data), "text"
    except ValueError as e:
        return JSONResponse({"error": str(e)}, status_code=400)
    specs, episode = await _import_specs(owner_ai_id, [], text)
    if not specs:
        return {"written": [], "episode": episode}
    written = await _write_view(owner_ai_id, 0, specs, episode, sense)
    retrieval.save_vectors()
    log.info("[import] 本体#%s 导入文件「%s」→ 写入 %s 颗种子", owner_ai_id, name, len(written))
    return {"written": written, "episode": episode}


@app.get("/api/seeds")
async def api_seeds(owner_ai_id: int, peer_ai_id: int = 0, limit: int = 100, seed_type: str = None,
                    user: dict = Depends(auth.get_current_user)):
    await auth.require_owner_ai(user, owner_ai_id)
    rows = await seeds_mod.active_seeds(owner_ai_id, peer_ai_id if peer_ai_id > 0 else None,
                                        seed_type=seed_type, limit=min(limit, 500))
    for r in rows:
        r["seed_date_fmt"] = seeds_mod.fmt_date_annotated(r.get("seed_date"))
    return {"seeds": rows}


@app.get("/api/working")
async def api_working(owner_ai_id: int, user: dict = Depends(auth.get_current_user)):
    """现量窗口：该本体尚未结生的临时记忆，按对话对端分组。

    这不是"记住了什么"，而是"此刻手里攒着什么、还没巩固成长期记忆"——
    控制台看它，等于看本体当下的工作记忆。
    """
    await auth.require_owner_ai(user, owner_ai_id)
    rows = await db.pool().query(
        "SELECT id, peer_ai_id, evidence, sense_id, seed_text, created_at, expire_at "
        "FROM alaya_seeds WHERE owner_ai_id=%s AND seed_type='sensory' AND state='temp' "
        "AND evidence IN (%s,%s) ORDER BY id",
        [int(owner_ai_id), working.MARK_SPEAKER, working.MARK_SELF])
    groups: dict[int, list] = {}
    for r in rows:
        groups.setdefault(int(r["peer_ai_id"] or 0), []).append({
            "id": int(r["id"]), "who": "本体" if r["evidence"] == working.MARK_SELF else "对方",
            "sense_id": str(r.get("sense_id") or ""), "text": str(r.get("seed_text") or ""),
            "at": str(r.get("created_at") or ""),
        })
    return {"total": len(rows), "groups": [{"peer_ai_id": k, "items": v} for k, v in groups.items()]}


@app.get("/api/selfmodel")
async def api_selfmodel(ai_id: int, peer_ai_id: int | None = None, user: dict = Depends(auth.get_current_user)):
    await auth.require_owner_ai(user, ai_id)
    # 两种情况（口径必须与引擎内部 _self_model_for 一致）：
    #  ① 指定 peer_ai_id：看「这个本体对某个对话对象」的我识——core 恒为 peer=0，relation/mood 取该对象那层。
    #  ② 不指定 peer_ai_id：看「这个本体的大脑实况」（控制台切谁看谁）——
    #     core（本体我）恒为 peer=0；relation/mood 取该本体最近演化过的那一层（当前最活跃的那段关系/心境）。
    #     ⚠️ 旧实现在 guide 模式下前端把「被观察本体」和「对话对象」搞反，切到如意却读到了用户本体的我识。
    # 前端要的是 {layer:{content,version,updated_at}}（旧实现只回裸 content → data.content 恒 undefined →
    # 界面永远"尚未成型"）。
    if peer_ai_id is None:
        rows = await db.pool().query(
            "SELECT layer, content, version, peer_ai_id, updated_at FROM alaya_self_model "
            "WHERE ai_id=%s ORDER BY updated_at DESC", [int(ai_id)])
        sm: dict = {}
        for r in rows:
            layer = str(r["layer"])
            # core 是本体恒定的那一层，只存在于 peer=0；其余层取最近更新的那条
            if layer == "core" and int(r["peer_ai_id"] or 0) != 0:
                continue
            sm.setdefault(layer, {
                "content": db.jload(r["content"], {}) or {},
                "version": int(r["version"] or 1),
                "updated_at": str(r["updated_at"] or ""),
            })
    else:
        rows = await db.pool().query(
            "SELECT layer, content, version, updated_at FROM alaya_self_model "
            "WHERE ai_id=%s AND (peer_ai_id=%s OR peer_ai_id=0) ORDER BY peer_ai_id DESC",
            [int(ai_id), int(peer_ai_id or 0)])
        sm = {}
        for r in rows:
            # 同层以「对话对象专属」优先（ORDER BY peer DESC 保证专属行先到；core 天生只有 peer=0）
            sm.setdefault(str(r["layer"]), {
                "content": db.jload(r["content"], {}) or {},
                "version": int(r["version"] or 1),
                "updated_at": str(r["updated_at"] or ""),
            })
    flat = {k: v["content"] for k, v in sm.items()}
    return {"self_model": sm, "prompt": selfmodel.build_self_prompt(flat)}


@app.get("/api/xixi_log")
async def api_xixi_log(limit: int = 100, user: dict = Depends(auth.get_current_user)):
    rows = await db.pool().query(
        "SELECT * FROM alaya_xixi_log ORDER BY id DESC LIMIT %s", [min(limit, 500)])
    return {"log": rows}


@app.get("/api/network")
async def api_network(owner_ai_id: int, peer_ai_id: int = 0, user: dict = Depends(auth.get_current_user)):
    """习气网络：概念节点 + 概念边 + 种子边（控制台可视化数据源）"""
    owner = int(owner_ai_id)
    if owner <= 0:
        return JSONResponse({"error": "owner_ai_id 必填"}, status_code=400)
    await auth.require_owner_ai(user, owner)
    concepts = await db.pool().query(
        "SELECT id, name, branch, weight, status FROM alaya_concepts WHERE owner_ai_id=%s ORDER BY weight DESC LIMIT 200",
        [owner])
    edges = await db.pool().query(
        "SELECT ce.src, ce.dst, ce.relation_type, ce.weight, sn.name AS src_name, dn.name AS dst_name "
        "FROM alaya_concept_edges ce "
        "JOIN alaya_concepts sn ON sn.id = ce.src "
        "JOIN alaya_concepts dn ON dn.id = ce.dst "
        "WHERE ce.owner_ai_id=%s ORDER BY ce.weight DESC LIMIT 200", [owner])
    seed_edge_n = await db.pool().query(
        "SELECT COUNT(*) AS n, COALESCE(SUM(weight),0) AS w FROM alaya_seed_edges WHERE owner_ai_id=%s", [owner])
    type_stat = await db.pool().query(
        "SELECT seed_type, COUNT(*) AS n, ROUND(SUM(weight),3) AS w FROM alaya_seeds "
        "WHERE owner_ai_id=%s AND state IN ('confirmed','temp') GROUP BY seed_type", [owner])
    return {
        "concepts": concepts,
        "edges": edges,
        "seed_edges": {"count": int(seed_edge_n[0]["n"]) if seed_edge_n else 0,
                       "weight": float(seed_edge_n[0]["w"]) if seed_edge_n else 0.0},
        "type_stat": type_stat,
    }


@app.post("/api/chat")
async def api_chat(body: dict, user: dict = Depends(auth.get_current_user)):
    """文本对话（控制台/联调用）：A=用户本人的AI本体，B=出场的对方AI本体——记忆双写两份。

    与自己的AI本体对话被拒绝：用户只能跟「其它」AI本体聊。"""
    chat_id = int(body.get("owner_ai_id") or 0)
    text_in = str(body.get("text") or "").strip()
    scene = str(body.get("scene") or "").strip()
    if chat_id <= 0 or (not text_in and not scene):
        return JSONResponse({"error": "owner_ai_id 与 text/scene 必填"}, status_code=400)
    await auth.require_owner_ai(user, chat_id)
    rows = await db.pool().query("SELECT id, name, user_id FROM xd_ais WHERE id=%s", [chat_id])
    if not rows:
        return JSONResponse({"error": "AI不存在"}, status_code=404)
    my_ai_id = await _user_main_ai(user["id"], fallback=0) or await auth._ensure_user_ai(int(user["id"]))
    if not my_ai_id:
        return JSONResponse({"error": "当前账号的AI本体不可用"}, status_code=400)
    if chat_id == int(my_ai_id):
        return JSONResponse({"error": "不能与自己的AI本体对话，请选择其它AI本体"}, status_code=403)
    chat_name = rows[0]["name"]
    login_name = str(user.get("nickname") or "")
    owner, peer = int(my_ai_id), chat_id              # 记忆归属A / 记忆归属B
    mode = "guide" if int(rows[0].get("user_id") or 0) == 0 else "persona"
    if mode == "guide":
        own_rows = await db.pool().query("SELECT name FROM xd_ais WHERE id=%s", [owner])
        owner_name = own_rows[0]["name"] if own_rows else "TA的本体"
    else:
        owner_name = chat_name
    speaker_name = login_name
    if not text_in:
        text_in = "（发来一张照片）"
    now_text = f"{text_in}\n（此刻我看到的画面：{scene}）" if scene else text_in
    await retrieval.load_vectors_once()
    system_prompt, recall = await _build_context(
        owner, peer, now_text, speaker_name=speaker_name, mode=mode,
        ai_name=chat_name, owner_name=owner_name,
        login_name=login_name if mode == "guide" else "",
        time_note=await _time_note(owner, peer))
    history = await _load_history(owner, peer)
    if scene:
        await _save_dialogue(owner, peer, "user", f"【画面】{scene}")
    await _save_dialogue(owner, peer, "user", text_in)
    messages = [{"role": "system", "content": system_prompt}, *history,
                {"role": "user", "content": now_text}]
    full = ""
    async for delta in llm_stream(messages, temperature=0.8, max_tokens=1024):
        full += delta
    await _save_dialogue(owner, peer, "assistant", full)
    mem = await _memory_pipeline(owner, peer, now_text, full,
                                 chat_name=chat_name, speaker_name=speaker_name,
                                 sense="photo" if scene else "voice",
                                 peer_distill=(mode != "guide"))
    return {"reply": full, "mode": mode, "owner_name": owner_name,
            "recalled": recall.get("seeds", []), "memory": mem}


HISTORY_PAGE = 20


@app.get("/api/history")
async def api_history(owner_ai_id: int, peer_ai_id: int = 0, before_id: int = 0, limit: int = HISTORY_PAGE,
                      user: dict = Depends(auth.get_current_user)):
    """对话记录：默认最近20条；before_id 懒加载更早的（前端滚到顶时续拉）"""
    await auth.require_owner_ai(user, owner_ai_id)
    owner, peer = int(owner_ai_id), int(peer_ai_id or 0)
    n = max(1, min(int(limit or HISTORY_PAGE), 50))
    sql = ("SELECT id, role, content FROM alaya_dialogues WHERE owner_ai_id=%s AND peer_ai_id=%s"
           + (" AND id<%s" if before_id > 0 else "") +
           " ORDER BY id DESC LIMIT %s")
    params = [owner, peer, before_id, n] if before_id > 0 else [owner, peer, n]
    rows = await db.pool().query(sql, params)
    rows.reverse()   # 正序返回（旧→新），前端直接按序渲染
    return {"messages": rows, "has_more": len(rows) >= n}


@app.websocket("/ws")
async def ws_endpoint(ws: WebSocket):
    await ws_converse(ws)


os.makedirs(PUBLIC_DIR, exist_ok=True)
app.mount("/", StaticFiles(directory=PUBLIC_DIR, html=True), name="public")


@app.middleware("http")
async def no_cache_html(request, call_next):
    """HTML 页面禁缓存——保证前端更新即刻生效，避免浏览器跑旧版JS打到错误路径"""
    resp = await call_next(request)
    if resp.headers.get("content-type", "").startswith("text/html"):
        resp.headers["Cache-Control"] = "no-cache"
    return resp


if __name__ == "__main__":
    uvicorn.run(app, host=str(settings.get("server.host", "0.0.0.0")),
                port=int(settings.get("server.port", 9890)))
